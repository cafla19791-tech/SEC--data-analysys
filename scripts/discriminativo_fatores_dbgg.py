#!/usr/bin/env python3
"""Discriminativo mensal dos fatores condicionantes da DBGG.

Fontes (Banco Central do Brasil):
- ftp/notaecon/facdbp.zip — série longa, fluxos mensais (dez/2006–nov/2018)
- Notas de estatísticas fiscais, Tabela 18 — vintages mensais (a partir de mai/2018),
  que revisam os meses recentes e estendem a série até o último dado publicado.

Em cada mês prevalece a publicação mais recente que o contém.
"""

from __future__ import annotations

import io
import re
import zipfile
from calendar import monthrange
from datetime import date
from pathlib import Path

import pandas as pd
import requests

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
FACDBP_URL = "https://www.bcb.gov.br/ftp/notaecon/facdbp.zip"
MENSAL_URL = (
    "https://www.bcb.gov.br/content/estatisticas/"
    "hist_estatisticasfiscais/{ym}_Tabelas_de_estatisticas_fiscais.xlsx"
)

MESES = {
    "janeiro": 1,
    "fevereiro": 2,
    "março": 3,
    "marco": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}

# rótulo normalizado -> campo
CAMPOS = [
    ("saldo", ("divida bruta do governo geral - saldo", "divida bruta do gov. geral - saldo")),
    ("variacao", ("variacao mensal",)),
    ("fatores", ("fatores condicionantes",)),
    ("necessidade", ("nec. de financiamento", "necessidades financ")),
    ("emissoes", ("emissoes liquidas",)),
    ("juros", ("juros nominais",)),
    ("ajuste", ("ajuste cambial",)),
    ("interna_cambio", ("interna indexada ao cambio",)),
    ("externa_met", ("divida externa - metodologico", "divida externa – metodologico")),
    ("externa_outros", ("divida externa - outros", "divida externa – outros")),
    ("reconhecimento", ("reconhecimento de dividas",)),
    ("privatizacoes", ("privatizacoes",)),
    ("efeito_pib", ("efeito do crescimento do pib", "efeito crescimento pib")),
    ("pib12", ("pib acumulado em doze meses", "pib ultimos 12")),
]


def _norm(texto: object) -> str:
    s = "" if texto is None else str(texto)
    s = s.lower().strip()
    s = (
        s.replace("á", "a")
        .replace("à", "a")
        .replace("ã", "a")
        .replace("â", "a")
        .replace("é", "e")
        .replace("ê", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ô", "o")
        .replace("õ", "o")
        .replace("ú", "u")
        .replace("ç", "c")
    )
    s = s.replace("–", "-").replace("—", "-")
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r"[:/0-9*]+$", "", s).strip()
    return s


def _campo(rotulo: object) -> str | None:
    n = _norm(rotulo)
    if not n or len(n) > 90:
        return None
    # título da planilha, não a linha de dados
    if n.startswith("divida bruta") and "fatores condicionantes" in n and "saldo" not in n and "variacao" not in n:
        return None
    # saldo aparece duas vezes (R$ e % do PIB); o parser trata o bloco único do xlsx
    for nome, chaves in CAMPOS:
        if any(n.startswith(c) or c in n for c in chaves):
            # 'variacao' não pode casar com o efeito do PIB
            if nome == "variacao" and "pib" in n:
                continue
            if nome == "saldo" and "variacao" in n:
                continue
            if nome == "fatores" and not n.startswith("fatores condicionantes"):
                continue
            return nome
    return None


def _num(v: object) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if not s or s in {"-", "–", "..."}:
        return None
    s = s.replace(".", "").replace(",", ".") if re.search(r",\d", s) else s
    try:
        return float(s)
    except ValueError:
        return None


def _mes_num(v: object) -> int | None:
    if v is None:
        return None
    return MESES.get(_norm(v))


def parse_facdbp(conteudo: bytes) -> list[dict]:
    """Série longa: bloco em R$ (linhas ~9-22) e bloco em % do PIB (linhas ~27-42)."""
    import xlrd

    wb = xlrd.open_workbook(file_contents=conteudo)
    sh = wb.sheet_by_name("Fluxos mensais")
    # anos na linha 4, meses na linha 6 (0-index)
    anos: dict[int, int] = {}
    ano_atual = None
    for c in range(sh.ncols):
        v = sh.cell_value(4, c)
        if isinstance(v, (int, float)) and v > 1900:
            ano_atual = int(v)
        if ano_atual:
            anos[c] = ano_atual

    # mapa campo -> linha no bloco de R$ e no bloco de %
    linhas_rs: dict[str, int] = {}
    linhas_pct: dict[str, int] = {}
    viu_pct = False
    for r in range(sh.nrows):
        campo = _campo(sh.cell_value(r, 0))
        if campo is None:
            continue
        rot = _norm(sh.cell_value(r, 0))
        if campo == "saldo" and "variacao" not in rot:
            # o segundo saldo (depois do PIB em R$) abre o bloco %
            if "saldo" in linhas_rs and r > 24:
                viu_pct = True
        if not viu_pct:
            linhas_rs.setdefault(campo, r)
        else:
            linhas_pct.setdefault(campo, r)

    regs = []
    for c in range(2, sh.ncols):
        mes = _mes_num(sh.cell_value(6, c))
        ano = anos.get(c)
        if not mes or not ano:
            continue
        reg = {"ano": ano, "mes": mes, "vintage": "facdbp-201812"}
        for campo, r in linhas_rs.items():
            if campo == "efeito_pib":
                continue
            reg[f"{campo}_rs"] = _num(sh.cell_value(r, c))
        for campo, r in linhas_pct.items():
            reg[f"{campo}_pib"] = _num(sh.cell_value(r, c))
        # efeito só existe em % do PIB
        if "efeito_pib" in linhas_pct:
            reg["efeito_pib_pib"] = _num(sh.cell_value(linhas_pct["efeito_pib"], c))
        regs.append(reg)
    return regs


def parse_tabela18(conteudo: bytes, vintage: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    if "Tabela 18" not in wb.sheetnames:
        wb.close()
        return []
    ws = wb["Tabela 18"]
    linhas = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(linhas) < 20:
        return []

    ano_atual = None
    anos: dict[int, int] = {}
    for c, v in enumerate(linhas[4]):
        if isinstance(v, (int, float)) and v > 1900:
            ano_atual = int(v)
        if ano_atual:
            anos[c] = ano_atual

    campos_linha: dict[str, int] = {}
    for i, row in enumerate(linhas):
        campo = _campo(row[0] if row else None)
        if campo:
            campos_linha.setdefault(campo, i)

    regs = []
    meses_row = linhas[6]
    for c, nome_mes in enumerate(meses_row):
        mes = _mes_num(nome_mes)
        if not mes:
            continue
        ano = anos.get(c)
        if not ano:
            continue
        reg = {"ano": ano, "mes": mes, "vintage": vintage}
        for campo, i in campos_linha.items():
            row = linhas[i]
            valor = _num(row[c]) if c < len(row) else None
            pct = _num(row[c + 1]) if c + 1 < len(row) else None
            if campo == "efeito_pib":
                reg["efeito_pib_pib"] = pct if pct is not None else valor
            elif campo == "pib12":
                reg["pib12_rs"] = valor
            else:
                reg[f"{campo}_rs"] = valor
                reg[f"{campo}_pib"] = pct
        regs.append(reg)
    return regs


def _ym_range(inicio: tuple[int, int], fim: tuple[int, int]) -> list[str]:
    y, m = inicio
    out = []
    while (y, m) <= fim:
        out.append(f"{y}{m:02d}")
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def baixar(url: str) -> bytes:
    import hashlib

    cache = Path("/tmp/bcb/cache")
    cache.mkdir(parents=True, exist_ok=True)
    dest = cache / hashlib.sha1(url.encode()).hexdigest()
    if dest.exists() and dest.stat().st_size > 1000:
        return dest.read_bytes()
    r = requests.get(url, timeout=120, headers={"User-Agent": "Mozilla/5.0"})
    r.raise_for_status()
    dest.write_bytes(r.content)
    return r.content


def consolidar(inicio: date, fim: date) -> pd.DataFrame:
    z = baixar(FACDBP_URL)
    with zipfile.ZipFile(io.BytesIO(z)) as arc:
        nome = next(n for n in arc.namelist() if n.lower().endswith(".xls"))
        base = parse_facdbp(arc.read(nome))

    # vintages mensais: publicação de mai/2018 (dado até abr/2018) até o mês seguinte ao fim
    pub_ini = (2018, 5)
    pub_fim_y, pub_fim_m = fim.year, fim.month + 1
    if pub_fim_m == 13:
        pub_fim_y, pub_fim_m = pub_fim_y + 1, 1
    vintages: list[dict] = []
    for ym in _ym_range(pub_ini, (pub_fim_y, pub_fim_m)):
        try:
            blob = baixar(MENSAL_URL.format(ym=ym))
        except requests.HTTPError:
            continue
        vintages.extend(parse_tabela18(blob, ym))

    # publicação mais recente ganha
    por_mes: dict[tuple[int, int], dict] = {}
    for reg in base:
        por_mes[(reg["ano"], reg["mes"])] = reg
    for reg in vintages:
        por_mes[(reg["ano"], reg["mes"])] = reg

    linhas = []
    chaves = sorted(k for k in por_mes if date(k[0], k[1], 1) >= date(inicio.year, inicio.month, 1) and date(k[0], k[1], 1) <= date(fim.year, fim.month, 1))
    # saldo de abertura = saldo do mês anterior (pode ser dez/2006)
    for ano, mes in chaves:
        reg = dict(por_mes[(ano, mes)])
        if mes == 1:
            prev = (ano - 1, 12)
        else:
            prev = (ano, mes - 1)
        reg["saldo_abertura_rs"] = por_mes.get(prev, {}).get("saldo_rs")
        reg["saldo_abertura_pib"] = por_mes.get(prev, {}).get("saldo_pib")
        reg["data"] = date(ano, mes, monthrange(ano, mes)[1])
        linhas.append(reg)

    df = pd.DataFrame(linhas).sort_values(["ano", "mes"]).reset_index(drop=True)
    return df


COLUNAS = [
    ("data", "Mês"),
    ("saldo_abertura_rs", "Dívida bruta — abertura (R$ milhões)"),
    ("saldo_abertura_pib", "Dívida bruta — abertura (% PIB)"),
    ("fatores_rs", "Fatores condicionantes (R$ milhões)"),
    ("fatores_pib", "Fatores condicionantes (% PIB)"),
    ("necessidade_rs", "Necessidade de financiamento da DBGG (R$ milhões)"),
    ("necessidade_pib", "Necessidade de financiamento da DBGG (% PIB)"),
    ("juros_rs", "Juros nominais (R$ milhões)"),
    ("juros_pib", "Juros nominais (% PIB)"),
    ("emissoes_rs", "Emissões líquidas (R$ milhões)"),
    ("emissoes_pib", "Emissões líquidas (% PIB)"),
    ("ajuste_rs", "Ajuste cambial (R$ milhões)"),
    ("ajuste_pib", "Ajuste cambial (% PIB)"),
    ("interna_cambio_rs", "Dívida interna indexada ao câmbio (R$ milhões)"),
    ("interna_cambio_pib", "Dívida interna indexada ao câmbio (% PIB)"),
    ("externa_met_rs", "Dívida externa — metodológico (R$ milhões)"),
    ("externa_met_pib", "Dívida externa — metodológico (% PIB)"),
    ("externa_outros_rs", "Dívida externa — outros ajustes (R$ milhões)"),
    ("externa_outros_pib", "Dívida externa — outros ajustes (% PIB)"),
    ("reconhecimento_rs", "Reconhecimento de dívidas (R$ milhões)"),
    ("reconhecimento_pib", "Reconhecimento de dívidas (% PIB)"),
    ("privatizacoes_rs", "Privatizações (R$ milhões)"),
    ("privatizacoes_pib", "Privatizações (% PIB)"),
    ("efeito_pib_pib", "Efeito do crescimento do PIB sobre a dívida (p.p. do PIB)"),
    ("variacao_rs", "Variação mensal da dívida bruta (R$ milhões)"),
    ("variacao_pib", "Variação mensal da dívida bruta (p.p. do PIB)"),
    ("saldo_rs", "Dívida bruta — fechamento (R$ milhões)"),
    ("saldo_pib", "Dívida bruta — fechamento (% PIB)"),
    ("pib12_rs", "PIB acumulado em 12 meses (R$ milhões)"),
    ("vintage", "Publicação BCB"),
]


def exportar(df: pd.DataFrame, destino: Path) -> None:
    tab = df.rename(columns={src: nome for src, nome in COLUNAS if src in df.columns})
    ordem = [nome for _, nome in COLUNAS if nome in tab.columns]
    tab = tab[ordem]
    tab["Mês"] = pd.to_datetime(tab["Mês"])
    destino.parent.mkdir(parents=True, exist_ok=True)
    csv_path = destino.with_suffix(".csv")
    tab.to_csv(csv_path, index=False, date_format="%Y-%m-%d", float_format="%.6f")

    with pd.ExcelWriter(destino, engine="openpyxl") as xl:
        tab.to_excel(xl, sheet_name="Mensal", index=False)
        notas = pd.DataFrame(
            {
                "Nota": [
                    "Dívida Bruta do Governo Geral (DBGG), metodologia vigente desde 2008, com série iniciada em dezembro de 2006.",
                    "Identidade em R$ milhões: fechamento = abertura + variação mensal.",
                    "Variação mensal (R$) = necessidade de financiamento + ajuste cambial + dívida externa (outros ajustes) + reconhecimento de dívidas + privatizações.",
                    "Necessidade de financiamento da DBGG = emissões líquidas + juros nominais. Não é o resultado primário nem a NFSP do setor público: emissões líquidas incluem o primário do governo geral e outras operações financeiras (reservas, bancos oficiais etc.).",
                    "Ajuste cambial = dívida interna indexada ao câmbio + dívida externa (ajuste metodológico).",
                    "Em % do PIB, a variação mensal da relação dívida/PIB = fatores condicionantes (% do PIB dos últimos 12 meses) + efeito do crescimento do PIB. O efeito do PIB não tem coluna em R$ milhões na Tabela 18 do BCB.",
                    "Fatores em % do PIB = (fluxo em R$ / PIB acumulado em 12 meses) × 100. Esse percentual não é a variação da relação dívida/PIB.",
                    "Para cada mês usa-se a nota de estatísticas fiscais mais recente que o publica. Meses anteriores a dez/2017 vêm da série histórica facdbp (atualização de dez/2018), porque as notas mensais só republicam cerca de cinco meses. Uma revisão posterior do PIB nominal pode alterar a razão dívida/PIB na Tabela 19 sem reescrever o mês na Tabela 18; os fluxos em R$ milhões não mudam por isso.",
                    "Fontes: Banco Central do Brasil, Tabela 18 das estatísticas fiscais e série histórica Facdbp (ftp/notaecon/facdbp.zip).",
                ]
            }
        )
        notas.to_excel(xl, sheet_name="Notas", index=False)

        ws = xl.book["Mensal"]
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "B2"
        ws.column_dimensions["A"].width = 14
        for col in ws.iter_cols(min_col=2, max_col=ws.max_column):
            ws.column_dimensions[col[0].column_letter].width = 22
        from openpyxl.styles import Alignment

        for cell in ws[1]:
            cell.alignment = Alignment(wrap_text=True, vertical="bottom")


def validar(df: pd.DataFrame) -> None:
    tol_rs = 0.05  # R$ milhões
    tol_pib = 0.002
    nec = df["emissoes_rs"] + df["juros_rs"]
    ajuste = df["interna_cambio_rs"] + df["externa_met_rs"]
    fatores = (
        df["necessidade_rs"]
        + df["ajuste_rs"]
        + df["externa_outros_rs"]
        + df["reconhecimento_rs"]
        + df["privatizacoes_rs"]
    )
    fecha = df["saldo_abertura_rs"] + df["variacao_rs"]
    var_pib = df["fatores_pib"] + df["efeito_pib_pib"]
    def _max(s: pd.Series) -> float:
        if s.isna().any():
            return float("nan")
        return float(s.abs().max())

    erros = {
        "necessidade": _max(nec - df["necessidade_rs"]),
        "ajuste": _max(ajuste - df["ajuste_rs"]),
        "fatores": _max(fatores - df["fatores_rs"]),
        "variacao_rs_vs_fatores": _max(df["variacao_rs"] - df["fatores_rs"]),
        "fechamento": _max(fecha - df["saldo_rs"]),
        "variacao_pib": _max(var_pib - df["variacao_pib"]),
    }
    print("checagens (máximo absoluto)")
    for k, v in erros.items():
        print(f"  {k}: {v:.6f}")
    if erros["necessidade"] > tol_rs or erros["ajuste"] > tol_rs or erros["fatores"] > tol_rs:
        raise SystemExit("identidade em R$ dos fatores não fechou")
    if erros["fechamento"] > 1.0:
        raise SystemExit("saldo de fechamento não fecha com a abertura")
    if erros["variacao_pib"] > tol_pib:
        raise SystemExit("variação em p.p. do PIB não fecha com fatores + efeito do PIB")
    if df[
        [
            "saldo_rs",
            "fatores_rs",
            "necessidade_rs",
            "juros_rs",
            "emissoes_rs",
            "ajuste_rs",
            "interna_cambio_rs",
            "externa_met_rs",
            "externa_outros_rs",
            "reconhecimento_rs",
            "privatizacoes_rs",
            "efeito_pib_pib",
            "saldo_abertura_rs",
        ]
    ].isna().any().any():
        raise SystemExit("há meses sem saldo, fatores, juros ou efeito do PIB")


def main() -> None:
    inicio = date(2007, 1, 1)
    fim = date(2026, 7, 1)
    df = consolidar(inicio, fim)
    esperado = (fim.year - inicio.year) * 12 + (fim.month - inicio.month) + 1
    print(f"meses: {len(df)} (esperado {esperado})")
    print(f"de {df.iloc[0]['ano']}-{df.iloc[0]['mes']:02d} a {df.iloc[-1]['ano']}-{df.iloc[-1]['mes']:02d}")
    validar(df)
    destino = OUT / "discriminativo_fatores_dbgg_2007_2026.xlsx"
    exportar(df, destino)
    print(destino)
    print(destino.with_suffix(".csv"))


if __name__ == "__main__":
    main()
