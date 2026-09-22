#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Discriminativo da evolução do lucro líquido e do ativo dos bancos (2002–2026).

Universo
--------
Conglomerados financeiros e instituições independentes do **segmento bancário**
(IFData, ``TipoInstituicao=2``). Cooperativas de crédito (TCB B3) e
conglomerados prudenciais ficam de fora, para não misturar sistemas nem
contar duas vezes o mesmo grupo.

Fonte
-----
Banco Central do Brasil, IFData — relatório Resumo (ativo total e lucro líquido):
https://dadosabertos.bcb.gov.br/dataset/ifdata---dados-selecionados-de-instituies-financeiras

Lucro líquido
-------------
No COSIF as contas de resultado acumulam no semestre e são zeradas em 30/06
e em 31/12. Por isso:

- data-base junho = lucro do 1º semestre;
- data-base dezembro = lucro do 2º semestre;
- lucro líquido do ano = soma dos dois semestres.

O ativo é estoque de 31/12. Banco que muda de código no meio do ano
(instituição avulsa → conglomerado) entra uma vez só. Quem deixa de
existir antes de dezembro não repete o ativo de junho — ele já está
no balanço de quem o absorveu. Em 2026 a última data-base publicada é
junho (ativo em 30/06 e lucro só do 1º semestre).

Uso::

  python3 scripts/discriminativo_lucros_ativos_bancos.py
  python3 scripts/discriminativo_lucros_ativos_bancos.py --saida-dir output
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import sys
import time
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Optional
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

ANO_MIN = 2002
ANO_MAX = 2026
TIPO_CONGLOMERADO_FINANCEIRO = 2
RELATORIO_RESUMO = "1"
OLINDA = "https://olinda.bcb.gov.br/olinda/servico/IFDATA/versao/v1/odata"
UA = "SEC-data-analysys/discriminativo-lucros-ativos-bancos"

# TCB bancário. B3S/B3C são cooperativas de crédito.
TCB_BANCO = {"B1", "B2", "B4"}
SEGMENTOS_BANCO = {
    "Banco Múltiplo",
    "Banco Comercial",
    "Banco de Investimento",
    "Banco Comercial Estrangeiro - Filial no país",
    "Banco de Câmbio",
    "Banco de Desenvolvimento",
    "Caixa Econômica",
    "Caixa Econômica Federal",
}
CONTROLE = {1: "Público", 2: "Privado nacional", 3: "Controle estrangeiro"}

# Abaixo disto a data-base ainda não foi publicada (resposta vazia da API).
MIN_INSTITUICOES = 30


def _sem_acento(texto: str) -> str:
    bruto = unicodedata.normalize("NFKD", texto or "")
    return "".join(c for c in bruto if not unicodedata.combining(c))


def normalizar_conta(nome: str) -> Optional[str]:
    """Mapeia o nome da coluna do Resumo para ``ativo`` ou ``lucro``."""
    base = _sem_acento(nome or "").split("(")[0].strip().lower()
    if base.startswith("ativo total"):
        return "ativo"
    if base.startswith("lucro liquido"):
        return "lucro"
    return None


def parse_saldo(valor: Any) -> Optional[float]:
    """Converte saldo IFData (texto pt-BR, número ou vazio) em float."""
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        if pd.isna(valor):
            return None
        return float(valor)
    texto = str(valor).strip()
    if texto == "" or texto.lower() in {"null", "none", "nan"}:
        return None
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def eh_banco(cad: dict) -> bool:
    """True para banco ou conglomerado financeiro bancário.

    Exclui cooperativa (TCB B3) e conglomerado prudencial (nome com
    ``PRUDENCIAL``), que duplicaria o conglomerado financeiro.
    """
    if not cad:
        return False
    tcb = str(cad.get("Tcb") or "").strip().upper()
    if tcb.startswith("B3"):
        return False
    nome = _sem_acento(str(cad.get("NomeInstituicao") or "")).upper()
    if "PRUDENCIAL" in nome:
        return False
    segmento = str(cad.get("SegmentoTb") or "").strip()
    if segmento in SEGMENTOS_BANCO:
        return True
    return tcb in TCB_BANCO


def controle_nome(tc: Any) -> str:
    try:
        return CONTROLE.get(int(tc), "")
    except (TypeError, ValueError):
        return ""


def _get_bytes(url: str, timeout: int = 180, tentativas: int = 5) -> bytes:
    last: Exception | None = None
    for i in range(tentativas):
        try:
            req = Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
            with urlopen(req, timeout=timeout) as resp:
                return resp.read()
        except (HTTPError, URLError, TimeoutError, OSError) as exc:
            last = exc
            time.sleep(min(2 ** i, 16))
    raise RuntimeError(f"Falha ao consultar {url}: {last}") from last


def _cache_path(cache_dir: Optional[Path], chave: str) -> Optional[Path]:
    if cache_dir is None:
        return None
    cache_dir.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(chave.encode("utf-8")).hexdigest()[:20]
    return cache_dir / f"{digest}.bin"


def _baixar(url: str, cache_dir: Optional[Path], chave: str) -> bytes:
    path = _cache_path(cache_dir, chave)
    if path is not None and path.exists() and path.stat().st_size > 0:
        return path.read_bytes()
    data = _get_bytes(url)
    if path is not None:
        path.write_bytes(data)
    return data


def url_valores(anomes: int) -> str:
    params = {
        "@AnoMes": str(anomes),
        "@TipoInstituicao": str(TIPO_CONGLOMERADO_FINANCEIRO),
        "@Relatorio": f"'{RELATORIO_RESUMO}'",
        "$format": "text/csv",
    }
    return (
        f"{OLINDA}/IfDataValores(AnoMes=@AnoMes,TipoInstituicao=@TipoInstituicao,Relatorio=@Relatorio)?"
        + urlencode(params)
    )


def url_cadastro(anomes: int) -> str:
    params = {"@AnoMes": str(anomes), "$format": "json"}
    return f"{OLINDA}/IfDataCadastro(AnoMes=@AnoMes)?" + urlencode(params)


def extrair_contas(rows: Iterable[dict]) -> dict[str, dict[str, float]]:
    """Agrega ativo e lucro por instituição, ignorando linhas repetidas."""
    vistos: dict[tuple[str, str], list[float]] = defaultdict(list)
    for row in rows:
        conta = normalizar_conta(str(row.get("NomeColuna") or ""))
        if conta is None:
            continue
        cod = str(row.get("CodInst") or "").strip()
        if not cod:
            continue
        saldo = parse_saldo(row.get("Saldo"))
        if saldo is None:
            continue
        vistos[(cod, conta)].append(saldo)

    saida: dict[str, dict[str, float]] = {}
    for (cod, conta), valores in vistos.items():
        # A API às vezes repete a mesma linha. Se divergir, fica o valor modal.
        escolhido = Counter(valores).most_common(1)[0][0]
        saida.setdefault(cod, {})[conta] = escolhido
    return saida


def _ler_csv(data: bytes) -> list[dict]:
    if not data or len(data) < 80:
        return []
    texto = data.decode("utf-8-sig", errors="replace")
    amostra = texto[:400]
    sep = ";" if amostra.count(";") > amostra.count(",") else ","
    return list(csv.DictReader(io.StringIO(texto), delimiter=sep))


def baixar_valores(anomes: int, cache_dir: Optional[Path] = None) -> dict[str, dict[str, float]]:
    data = _baixar(url_valores(anomes), cache_dir, f"valores-{anomes}")
    contas = extrair_contas(_ler_csv(data))
    if len(contas) < MIN_INSTITUICOES:
        return {}
    return contas


def baixar_cadastro(anomes: int, cache_dir: Optional[Path] = None) -> dict[str, dict]:
    data = _baixar(url_cadastro(anomes), cache_dir, f"cadastro-{anomes}")
    payload = json.loads(data.decode("utf-8"))
    registros = payload.get("value") or []
    cad: dict[str, dict] = {}
    for reg in registros:
        cod = str(reg.get("CodInst") or "").strip()
        if cod:
            cad[cod] = reg
    return cad


def _cadastro_do_banco(
    cod: str,
    cadastros: dict[int, dict[str, dict]],
    meses: list[int],
) -> dict:
    for anomes in meses:
        reg = cadastros.get(anomes, {}).get(cod)
        if reg:
            return reg
    return {}


def _conglomerado(reg: dict) -> str:
    cong = str((reg or {}).get("CodConglomeradoFinanceiro") or "").strip()
    if cong.lower() in {"", "none", "null"}:
        return ""
    return cong


def _grupos_do_ano(
    codigos: set[str],
    cadastros: dict[int, dict[str, dict]],
    meses: list[int],
) -> dict[str, list[str]]:
    """Agrupa a instituição no conglomerado quando os dois aparecem no mesmo ano.

    Isso acontece em migração de código (a Caixa, em 2021, saiu do CNPJ
    avulso no 1º semestre e passou ao conglomerado no 2º). Sem o grupo,
    o ativo de dezembro e o de junho seriam somados duas vezes.
    """
    grupos: dict[str, list[str]] = defaultdict(list)
    for cod in sorted(codigos):
        cong = _conglomerado(_cadastro_do_banco(cod, cadastros, meses))
        chave = cong if cong and cong != cod and cong in codigos else cod
        if cod not in grupos[chave]:
            grupos[chave].append(cod)
    for chave, membros in grupos.items():
        if chave not in membros:
            membros.insert(0, chave)
    return grupos


def _primeiro_valor(
    mapa: dict[str, dict[str, float]],
    membros: list[str],
    preferido: str,
    campo: str,
) -> Optional[float]:
    bloco = mapa.get(preferido) or {}
    if bloco.get(campo) is not None:
        return bloco[campo]
    for cod in membros:
        valor = (mapa.get(cod) or {}).get(campo)
        if valor is not None:
            return valor
    return None


def montar_discriminativo(
    por_anomes: dict[int, dict[str, dict[str, float]]],
    cadastros: dict[int, dict[str, dict]],
    *,
    ano_min: int = ANO_MIN,
    ano_max: int = ANO_MAX,
) -> pd.DataFrame:
    """Uma linha por banco e ano, com ativo de fechamento e lucro do exercício."""
    linhas: list[dict] = []
    for ano in range(ano_min, ano_max + 1):
        jun_id = ano * 100 + 6
        dez_id = ano * 100 + 12
        junho = por_anomes.get(jun_id) or {}
        dezembro = por_anomes.get(dez_id) or {}
        tem_dezembro = bool(dezembro)
        meses_presentes = [
            ano * 100 + mes
            for mes in (12, 9, 6, 3)
            if por_anomes.get(ano * 100 + mes)
        ]
        if not meses_presentes:
            continue
        codigos = set(junho) | set(dezembro)
        for anomes in meses_presentes:
            codigos |= set(por_anomes.get(anomes) or {})
        codigos_banco = {
            cod
            for cod in codigos
            if eh_banco(_cadastro_do_banco(cod, cadastros, meses_presentes))
        }
        grupos = _grupos_do_ano(codigos_banco, cadastros, meses_presentes)
        for chave, membros in grupos.items():
            candidatos = [
                (cod, _cadastro_do_banco(cod, cadastros, meses_presentes))
                for cod in membros
            ]
            candidatos = [(cod, reg) for cod, reg in candidatos if eh_banco(reg)]
            if not candidatos:
                continue
            cod, cad = next(((c, r) for c, r in candidatos if c == chave), candidatos[0])
            if tem_dezembro:
                ativo = _primeiro_valor(dezembro, membros, chave, "ativo")
                data_base_ativo = dez_id if ativo is not None else None
            else:
                ativo = None
                data_base_ativo = None
                for anomes in meses_presentes:
                    ativo = _primeiro_valor(por_anomes[anomes], membros, chave, "ativo")
                    if ativo is not None:
                        data_base_ativo = anomes
                        break
            lucro_1s = _primeiro_valor(junho, membros, chave, "lucro")
            lucro_2s = _primeiro_valor(dezembro, membros, chave, "lucro") if tem_dezembro else None
            if tem_dezembro:
                if lucro_1s is None and lucro_2s is None:
                    lucro = None
                else:
                    lucro = (lucro_1s or 0.0) + (lucro_2s or 0.0)
                completo = lucro_1s is not None and lucro_2s is not None
            else:
                lucro = lucro_1s
                completo = False
            if (ativo is None or ativo == 0.0) and (lucro is None or lucro == 0.0):
                continue
            linhas.append(
                {
                    "ano": ano,
                    "cod_inst": cod,
                    "nome": str(cad.get("NomeInstituicao") or "").strip(),
                    "tcb": str(cad.get("Tcb") or "").strip(),
                    "segmento": str(cad.get("SegmentoTb") or "").strip(),
                    "uf": str(cad.get("Uf") or "").strip(),
                    "controle": controle_nome(cad.get("Tc")),
                    "sr": str(cad.get("Sr") or "").strip(),
                    "ativo": ativo,
                    "data_base_ativo": data_base_ativo,
                    "lucro_1s": lucro_1s,
                    "lucro_2s": lucro_2s,
                    "lucro_liquido": lucro,
                    "exercicio_completo": bool(completo),
                    "ano_com_dezembro": tem_dezembro,
                }
            )
    if not linhas:
        return pd.DataFrame(
            columns=[
                "ano",
                "cod_inst",
                "nome",
                "tcb",
                "segmento",
                "uf",
                "controle",
                "sr",
                "ativo",
                "data_base_ativo",
                "lucro_1s",
                "lucro_2s",
                "lucro_liquido",
                "exercicio_completo",
                "ano_com_dezembro",
            ]
        )
    df = pd.DataFrame(linhas)
    df["lucro_sobre_ativo"] = df["lucro_liquido"] / df["ativo"].where(df["ativo"] != 0)
    df = df.sort_values(["ano", "ativo"], ascending=[True, False], na_position="last")
    return df.reset_index(drop=True)


def montar_evolucao(disc: pd.DataFrame) -> pd.DataFrame:
    """Soma anual do ativo e do lucro líquido do conjunto de bancos."""
    if disc.empty:
        return pd.DataFrame()
    rows = []
    for ano, g in disc.groupby("ano"):
        ativo = float(g["ativo"].sum(min_count=1))
        lucro = float(g["lucro_liquido"].sum(min_count=1))
        bases = [int(v) for v in g["data_base_ativo"].dropna().tolist()]
        base = Counter(bases).most_common(1)[0][0] if bases else None
        completo = bool(g["ano_com_dezembro"].iloc[0])
        rows.append(
            {
                "ano": int(ano),
                "n_bancos": int(g["ativo"].notna().sum()),
                "n_exercicio_completo": int(g["exercicio_completo"].sum()),
                "ativo": ativo,
                "ativo_r_bilhoes": ativo / 1e9,
                "data_base_ativo": base,
                "lucro_1s": float(g["lucro_1s"].sum(min_count=1)),
                "lucro_2s": float(g["lucro_2s"].sum(min_count=1)) if completo else None,
                "lucro_liquido": lucro,
                "lucro_r_bilhoes": lucro / 1e9,
                "exercicio_completo": completo,
                "lucro_sobre_ativo": (lucro / ativo) if ativo else None,
            }
        )
    out = pd.DataFrame(rows).sort_values("ano").reset_index(drop=True)
    out["var_ativo"] = out["ativo"].pct_change()
    var_lucro = []
    for i, row in out.iterrows():
        if i == 0 or not bool(row["exercicio_completo"]) or not bool(out.loc[i - 1, "exercicio_completo"]):
            var_lucro.append(None)
            continue
        anterior = out.loc[i - 1, "lucro_liquido"]
        if anterior in (0, None) or pd.isna(anterior):
            var_lucro.append(None)
        else:
            var_lucro.append(row["lucro_liquido"] / anterior - 1)
    out["var_lucro"] = var_lucro
    base_ativo = out.loc[out["exercicio_completo"], "ativo"]
    if not base_ativo.empty and base_ativo.iloc[0]:
        ancora = float(base_ativo.iloc[0])
        out["indice_ativo_ano_base"] = out["ativo"] / ancora * 100
    else:
        out["indice_ativo_ano_base"] = None
    return out


def matriz_por_ano(disc: pd.DataFrame, coluna: str, em_bilhoes: bool = True) -> pd.DataFrame:
    """Bancos nas linhas e anos nas colunas, ordenados pelo ativo mais recente."""
    if disc.empty:
        return pd.DataFrame()
    ultimo = int(disc["ano"].max())
    ordem = (
        disc.loc[disc["ano"] == ultimo]
        .sort_values("ativo", ascending=False)["cod_inst"]
        .tolist()
    )
    nomes = (
        disc.sort_values("ano")
        .groupby("cod_inst")["nome"]
        .last()
        .to_dict()
    )
    amplo = disc.pivot_table(index="cod_inst", columns="ano", values=coluna, aggfunc="sum")
    if em_bilhoes:
        amplo = amplo / 1e9
    resto = [c for c in amplo.index if c not in set(ordem)]
    amplo = amplo.reindex(ordem + resto)
    amplo.insert(0, "nome", [nomes.get(c, "") for c in amplo.index])
    amplo.index.name = "cod_inst"
    return amplo


def _fmt_num(valor: Any, casas: int = 1) -> str:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return "—"
    texto = f"{float(valor):,.{casas}f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def _fmt_pct(valor: Any) -> str:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return "—"
    return _fmt_num(float(valor) * 100, 1) + "%"


def _nome_curto(nome: str, limite: int = 42) -> str:
    nome = " ".join(str(nome or "").split())
    if len(nome) <= limite:
        return nome
    return nome[: limite - 1] + "…"


def escrever_markdown(
    evolucao: pd.DataFrame,
    disc: pd.DataFrame,
    path: Path,
    gerado_em: str,
) -> None:
    linhas = [
        "# Discriminativo — lucro líquido e ativo dos bancos (2002–2026)",
        "",
        f"**Gerado em:** {gerado_em}",
        "",
        "Evolução anual do **ativo total** e do **lucro líquido** dos bancos,",
        "em conglomerado financeiro (instituição independente quando não há conglomerado).",
        "Cooperativas de crédito ficam fora deste recorte.",
        "",
        "Fonte: [IFData](https://www3.bcb.gov.br/ifdata/) do Banco Central do Brasil,",
        "relatório Resumo, `TipoInstituicao = 2` (conglomerados financeiros e instituições independentes).",
        "Valores em reais correntes.",
        "",
        "Metodologia:",
        "",
        "- Ativo: estoque de 31/12. Código que migra no meio do ano (ex.: Caixa em 2021)",
        "  não é somado duas vezes. Em 2026, estoque de 30/06, última base publicada.",
        "- Lucro líquido do ano: soma do 1º semestre (junho) e do 2º semestre (dezembro).",
        "  No COSIF o resultado é encerrado em 30/06 e em 31/12; a data-base de dezembro",
        "  sozinha é só o segundo semestre.",
        "- 2026 está incompleto: ativo em 30/06/2026 e lucro apenas do 1º semestre.",
        "",
        "| Ano | Bancos | Ativo (R$ bi) | Lucro líquido (R$ bi) | Lucro/Ativo | Δ ativo | Δ lucro |",
        "|----:|-------:|---------------:|----------------------:|------------:|--------:|--------:|",
    ]
    for row in evolucao.itertuples(index=False):
        lucro_lbl = _fmt_num(row.lucro_r_bilhoes, 1)
        if not row.exercicio_completo:
            lucro_lbl += " †"
        linhas.append(
            f"| {int(row.ano)} | {int(row.n_bancos)} | "
            f"{_fmt_num(row.ativo_r_bilhoes, 1)} | {lucro_lbl} | "
            f"{_fmt_pct(row.lucro_sobre_ativo)} | {_fmt_pct(row.var_ativo)} | "
            f"{_fmt_pct(row.var_lucro)} |"
        )
    linhas.extend(
        [
            "",
            "† 2026: lucro do 1º semestre e ativo em junho. A variação do ativo compara junho/2026 com dezembro/2025.",
            "",
        ]
    )

    completos = evolucao.loc[evolucao["exercicio_completo"]]
    if not completos.empty and not disc.empty:
        ano_fim = int(completos["ano"].max())
        ano_ini = int(completos["ano"].min())
        topo = disc.loc[disc["ano"] == ano_fim].nlargest(15, "ativo")
        ini = disc.loc[disc["ano"] == ano_ini].set_index("cod_inst")
        parcial = disc.loc[disc["ano"] == ANO_MAX].set_index("cod_inst")
        ini_por_nome = (
            disc.loc[disc["ano"] == ano_ini]
            .assign(_nome=lambda f: f["nome"].map(lambda n: _sem_acento(str(n)).upper()))
            .sort_values("ativo")
            .groupby("_nome", as_index=True)
            .tail(1)
            .set_index("_nome")
        )
        linhas.extend(
            [
                f"## Maiores bancos em {ano_fim}",
                "",
                "Ativo e lucro líquido em R$ bilhões. O código é o do conglomerado financeiro no IFData.",
                "",
                f"| Banco | Ativo {ano_ini} | Ativo {ano_fim} | Lucro {ano_ini} | Lucro {ano_fim} | Ativo {ANO_MAX} † | Lucro {ANO_MAX} † |",
                "|---|---:|---:|---:|---:|---:|---:|",
            ]
        )
        for row in topo.itertuples(index=False):
            cod = row.cod_inst
            nome_chave = _sem_acento(str(row.nome)).upper()
            if cod in ini.index:
                a0 = ini.at[cod, "ativo"] / 1e9
                l0 = ini.at[cod, "lucro_liquido"] / 1e9
            elif nome_chave in ini_por_nome.index:
                a0 = ini_por_nome.at[nome_chave, "ativo"] / 1e9
                l0 = ini_por_nome.at[nome_chave, "lucro_liquido"] / 1e9
            else:
                a0 = None
                l0 = None
            a26 = parcial.at[cod, "ativo"] / 1e9 if cod in parcial.index else None
            l26 = parcial.at[cod, "lucro_liquido"] / 1e9 if cod in parcial.index else None
            linhas.append(
                f"| {_nome_curto(row.nome)} | {_fmt_num(a0, 1)} | "
                f"{_fmt_num(row.ativo / 1e9, 1)} | {_fmt_num(l0, 1)} | "
                f"{_fmt_num(row.lucro_liquido / 1e9, 1)} | {_fmt_num(a26, 1)} | "
                f"{_fmt_num(l26, 1)} |"
            )
        linhas.append("")

    if not evolucao.empty:
        ultimo_cheio = evolucao.loc[evolucao["exercicio_completo"]].tail(1)
        if not ultimo_cheio.empty:
            r = ultimo_cheio.iloc[0]
            linhas.append(
                f"Em {int(r.ano)} os bancos somavam {_fmt_num(r.ativo_r_bilhoes, 1)} bilhões de reais "
                f"de ativo e {_fmt_num(r.lucro_r_bilhoes, 1)} bilhões de lucro líquido "
                f"({_fmt_pct(r.lucro_sobre_ativo)} do ativo), em {int(r.n_bancos)} instituições."
            )
        primeiro = evolucao.iloc[0]
        linhas.append(
            f"Em {int(primeiro.ano)} o ativo era {_fmt_num(primeiro.ativo_r_bilhoes, 1)} bilhões "
            f"e o lucro líquido {_fmt_num(primeiro.lucro_r_bilhoes, 1)} bilhões."
        )
        linhas.append("")
    path.write_text("\n".join(linhas) + "\n", encoding="utf-8")


def escrever_excel(evolucao: pd.DataFrame, disc: pd.DataFrame, path: Path) -> None:
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ativo = matriz_por_ano(disc, "ativo", em_bilhoes=True)
    lucro = matriz_por_ano(disc, "lucro_liquido", em_bilhoes=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        evolucao.to_excel(writer, sheet_name="Evolucao", index=False)
        disc.to_excel(writer, sheet_name="Discriminativo", index=False)
        ativo.to_excel(writer, sheet_name="Ativo_R$bi")
        lucro.to_excel(writer, sheet_name="Lucro_R$bi")
        wb = writer.book
        ws = wb["Evolucao"]
        fill = PatternFill("solid", fgColor="1F4E79")
        font = Font(color="FFFFFF", bold=True)
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(1, col)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18
        ws.column_dimensions["A"].width = 10

        # Gráfico na aba Evolucao: ativo (barras) e lucro (linha), em R$ bi.
        if len(evolucao) >= 2:
            barras = BarChart()
            barras.type = "col"
            barras.title = "Ativo e lucro líquido dos bancos"
            barras.y_axis.title = "Ativo (R$ bi)"
            data_ativo = Reference(ws, min_col=5, min_row=1, max_row=1 + len(evolucao))
            cats = Reference(ws, min_col=1, min_row=2, max_row=1 + len(evolucao))
            barras.add_data(data_ativo, titles_from_data=True)
            barras.set_categories(cats)
            barras.shape = 4
            barras.y_axis.axId = 100
            linha = LineChart()
            linha.y_axis.axId = 200
            linha.y_axis.title = "Lucro líquido (R$ bi)"
            data_lucro = Reference(ws, min_col=10, min_row=1, max_row=1 + len(evolucao))
            linha.add_data(data_lucro, titles_from_data=True)
            linha.y_axis.crosses = "max"
            barras.y_axis.crosses = "min"
            barras += linha
            barras.y_axis.crosses = "min"
            barras.width = 22
            barras.height = 10
            ws.add_chart(barras, "A28")

        nota = wb.create_sheet("Nota", 0)
        nota["A1"] = (
            "Discriminativo do ativo total e do lucro líquido dos bancos, 2002–2026. "
            "Universo: conglomerados financeiros e instituições independentes do segmento "
            "bancário (IFData, TipoInstituicao=2, TCB B1/B2/B4 ou segmento bancário). "
            "Cooperativas de crédito e conglomerados prudenciais não entram. "
            "Ativo é o estoque de 31/12 (em 2026, 30/06). "
            "Lucro líquido anual = 1º semestre (data-base junho) + 2º semestre (data-base dezembro), "
            "porque o COSIF encerra as contas de resultado ao fim de cada semestre. "
            "Em 2026 o 2º semestre ainda não foi publicado: o lucro é só o 1º semestre. "
            "Abas Ativo_R$bi e Lucro_R$bi estão em R$ bilhões. As demais colunas monetárias "
            "da aba Discriminativo estão em reais."
        )
        nota["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        nota.column_dimensions["A"].width = 120
        nota.row_dimensions[1].height = 90


def grafico_evolucao(evolucao: pd.DataFrame, path: Path) -> None:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax1 = plt.subplots(figsize=(11, 5.5))
    anos = evolucao["ano"].astype(int)
    ax1.bar(anos, evolucao["ativo_r_bilhoes"], color="#1F4E79", width=0.7, label="Ativo")
    ax1.set_ylabel("Ativo (R$ bilhões)")
    ax1.set_xlabel("Ano")
    ax2 = ax1.twinx()
    ax2.plot(
        anos,
        evolucao["lucro_r_bilhoes"],
        color="#C65911",
        marker="o",
        linewidth=2,
        label="Lucro líquido",
    )
    ax2.set_ylabel("Lucro líquido (R$ bilhões)")
    ax1.set_title(
        "Bancos — ativo e lucro líquido (2002–2026)\n"
        "2026: ativo em junho e lucro apenas do 1º semestre"
    )
    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, loc="upper left")
    ax1.set_xticks(list(anos[::2]))
    fig.tight_layout()
    fig.savefig(path, dpi=140)
    plt.close(fig)


def processar(
    saida_dir: Path,
    *,
    ano_min: int = ANO_MIN,
    ano_max: int = ANO_MAX,
    cache_dir: Optional[Path] = None,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    saida_dir.mkdir(parents=True, exist_ok=True)
    por_anomes: dict[int, dict[str, dict[str, float]]] = {}
    cadastros: dict[int, dict[str, dict]] = {}
    for ano in range(ano_min, ano_max + 1):
        for mes in (6, 12):
            anomes = ano * 100 + mes
            print(f"[INFO] IFData {anomes} …", flush=True)
            try:
                valores = baixar_valores(anomes, cache_dir)
            except Exception as exc:
                print(f"  sem valores {anomes}: {exc}", flush=True)
                valores = {}
            if not valores:
                print(f"  {anomes} ainda não publicado", flush=True)
                continue
            por_anomes[anomes] = valores
            try:
                cadastros[anomes] = baixar_cadastro(anomes, cache_dir)
            except Exception as exc:
                raise RuntimeError(f"Cadastro {anomes} indisponível: {exc}") from exc
            print(
                f"  {len(valores):,} instituições no Resumo, "
                f"{len(cadastros[anomes]):,} no cadastro",
                flush=True,
            )

    disc = montar_discriminativo(por_anomes, cadastros, ano_min=ano_min, ano_max=ano_max)
    evolucao = montar_evolucao(disc)
    if disc.empty:
        raise RuntimeError("Nenhum banco classificado no período.")

    base = saida_dir / "discriminativo_lucros_ativos_bancos_2002_2026"
    disc.to_csv(base.with_suffix(".csv"), index=False, float_format="%.2f")
    evolucao.to_csv(
        saida_dir / "evolucao_lucros_ativos_bancos_2002_2026.csv",
        index=False,
        float_format="%.6f",
    )
    xlsx = base.with_suffix(".xlsx")
    escrever_excel(evolucao, disc, xlsx)
    gerado = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    escrever_markdown(evolucao, disc, base.with_suffix(".md"), gerado)
    grafico_evolucao(evolucao, base.with_suffix(".png"))
    print(f"[OK] {base.with_suffix('.csv')}")
    print(f"[OK] {xlsx}")
    print(f"[OK] {base.with_suffix('.md')}")
    print(f"[OK] {base.with_suffix('.png')}")
    return disc, evolucao


def parse_args(argv: Optional[list[str]] = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--saida-dir", type=Path, default=ROOT / "output")
    p.add_argument("--ano-min", type=int, default=ANO_MIN)
    p.add_argument("--ano-max", type=int, default=ANO_MAX)
    p.add_argument(
        "--cache-dir",
        type=Path,
        default=None,
        help="Pasta para reutilizar downloads do IFData",
    )
    return p.parse_args(argv)


def main(argv: Optional[list[str]] = None) -> int:
    args = parse_args(argv)
    try:
        disc, evolucao = processar(
            args.saida_dir,
            ano_min=args.ano_min,
            ano_max=args.ano_max,
            cache_dir=args.cache_dir,
        )
    except Exception as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        return 1
    print(evolucao[["ano", "n_bancos", "ativo_r_bilhoes", "lucro_r_bilhoes"]].to_string(index=False))
    cheio = evolucao.loc[evolucao["ano"] == 2024]
    if not cheio.empty:
        ativo = float(cheio["ativo_r_bilhoes"].iloc[0])
        if not 5_000 <= ativo <= 30_000:
            print(
                f"ERRO: ativo de 2024 fora da faixa esperada do sistema bancário ({ativo:.0f} bi).",
                file=sys.stderr,
            )
            return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
