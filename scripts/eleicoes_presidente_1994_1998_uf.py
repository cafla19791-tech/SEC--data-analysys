#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Votos válidos dos dois candidatos à Presidência mais votados, 1994 e 1998.

1º turno em cada uma das 27 Unidades da Federação. Ambas as eleições
encerraram-se no primeiro turno (FHC × Lula).

Fonte dos microdados por UF: Election Resources (tabelas montadas a partir
da totalização do TSE). Conferência: páginas TSE 1994 e totais nacionais
publicados pelo Tribunal.

Uso::

  PYTHONPATH=. python3 scripts/eleicoes_presidente_1994_1998_uf.py
"""

from __future__ import annotations

import argparse
import html
import json
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

import pandas as pd
import requests
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

CACHE = ROOT / "data" / "eleicoes_presidente_1994_1998_uf.json"
URL = "http://electionresources.org/br/president.php"
UA = "SEC-data-analysys/eleicoes-presidente-1994-1998"
ANOS = (1994, 1998)
DATA = {1994: "03/10/1994", 1998: "04/10/1998"}

# Totais nacionais (TSE / Election Resources).
TSE_BRASIL = {
    1994: {"fhc": 34_364_961, "lula": 17_122_127, "vv": 63_312_331},
    1998: {"fhc": 35_936_382, "lula": 21_475_211, "vv": 67_722_303},
}

UFS = [
    "AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT", "MS",
    "MG", "PA", "PB", "PR", "PE", "PI", "RJ", "RN", "RS", "RO", "RR", "SC",
    "SP", "SE", "TO",
]

NOMES = {
    "AC": "Acre",
    "AL": "Alagoas",
    "AP": "Amapá",
    "AM": "Amazonas",
    "BA": "Bahia",
    "CE": "Ceará",
    "DF": "Distrito Federal",
    "ES": "Espírito Santo",
    "GO": "Goiás",
    "MA": "Maranhão",
    "MT": "Mato Grosso",
    "MS": "Mato Grosso do Sul",
    "MG": "Minas Gerais",
    "PA": "Pará",
    "PB": "Paraíba",
    "PR": "Paraná",
    "PE": "Pernambuco",
    "PI": "Piauí",
    "RJ": "Rio de Janeiro",
    "RN": "Rio Grande do Norte",
    "RS": "Rio Grande do Sul",
    "RO": "Rondônia",
    "RR": "Roraima",
    "SC": "Santa Catarina",
    "SP": "São Paulo",
    "SE": "Sergipe",
    "TO": "Tocantins",
    "ZZ": "Exterior",
    "BR": "Brasil",
}

REGIAO = {
    "AC": "Norte",
    "AL": "Nordeste",
    "AP": "Norte",
    "AM": "Norte",
    "BA": "Nordeste",
    "CE": "Nordeste",
    "DF": "Centro-Oeste",
    "ES": "Sudeste",
    "GO": "Centro-Oeste",
    "MA": "Nordeste",
    "MT": "Centro-Oeste",
    "MS": "Centro-Oeste",
    "MG": "Sudeste",
    "PA": "Norte",
    "PB": "Nordeste",
    "PR": "Sul",
    "PE": "Nordeste",
    "PI": "Nordeste",
    "RJ": "Sudeste",
    "RN": "Nordeste",
    "RS": "Sul",
    "RO": "Norte",
    "RR": "Norte",
    "SC": "Sul",
    "SP": "Sudeste",
    "SE": "Nordeste",
    "TO": "Norte",
    "ZZ": "Exterior",
    "BR": "Brasil",
}

META_LABELS = {
    "Registered Electors",
    "Voters",
    "Blank Votes",
    "Invalid Votes",
    "Valid Votes",
    "Candidate",
}

COLUNAS = [
    "uf",
    "unidade",
    "regiao",
    "y1994_c1_votos",
    "y1994_c1_pct",
    "y1994_c2_votos",
    "y1994_c2_pct",
    "y1994_validos",
    "y1994_vencedor",
    "y1998_c1_votos",
    "y1998_c1_pct",
    "y1998_c2_votos",
    "y1998_c2_pct",
    "y1998_validos",
    "y1998_vencedor",
]


@dataclass(frozen=True)
class Candidato:
    chave: str
    nome: str
    votos_brasil: int
    pct_brasil: float


def parse_int(texto: str) -> int:
    return int(re.sub(r"[^0-9]", "", str(texto)))


def nome_curto(nome_completo: str) -> str:
    base = re.sub(r"\s*\(.*\)\s*$", "", nome_completo).strip()
    if "Fernando Henrique" in base:
        return "FHC"
    if "Lula" in base:
        return "Lula"
    if "Ciro" in base:
        return "Ciro Gomes"
    if "Enéas" in base or "Eneas" in base:
        return "Enéas"
    return base


def chave_candidato(nome_completo: str) -> str:
    if "Fernando Henrique" in nome_completo:
        return "fhc"
    if "Lula" in nome_completo:
        return "lula"
    return nome_curto(nome_completo).lower()


def parse_pagina(texto: str) -> dict[str, Any]:
    title = re.search(r"Presidential Election Results - ([^<]+)", texto)
    unidade = html.unescape(title.group(1)).strip() if title else None
    rows: list[tuple[str, int]] = []
    for m in re.finditer(
        r"<TD ALIGN=LEFT NOWRAP>&nbsp;([^<]+)&nbsp;</TD>\s*<TD>&nbsp;&nbsp;([^<]+)</TD>",
        texto,
    ):
        nome = html.unescape(m.group(1)).strip()
        votos = parse_int(m.group(2))
        rows.append((nome, votos))
    meta = {n: v for n, v in rows if n in META_LABELS}
    cands = [{"nome": n, "votos": v} for n, v in rows if n not in META_LABELS]
    return {
        "unidade": unidade,
        "eleitores": meta.get("Registered Electors"),
        "comparecimento": meta.get("Voters"),
        "brancos": meta.get("Blank Votes"),
        "nulos": meta.get("Invalid Votes"),
        "validos": meta.get("Valid Votes"),
        "candidatos": cands,
    }


def votos_de(cands: list[dict], chave: str) -> int:
    for item in cands:
        if chave_candidato(item["nome"]) == chave:
            return int(item["votos"])
    raise KeyError(chave)


def dois_mais_votados(payload_br: dict) -> tuple[Candidato, Candidato]:
    ordenados = sorted(payload_br["candidatos"], key=lambda c: int(c["votos"]), reverse=True)
    if len(ordenados) < 2:
        raise ValueError("Sem dois candidatos à Presidência.")
    validos = int(payload_br["validos"])
    top: list[Candidato] = []
    for item in ordenados[:2]:
        votos = int(item["votos"])
        top.append(
            Candidato(
                chave=chave_candidato(item["nome"]),
                nome=nome_curto(item["nome"]),
                votos_brasil=votos,
                pct_brasil=round(100.0 * votos / validos, 2) if validos else 0.0,
            )
        )
    return top[0], top[1]


def vencedor_local(cands: list[dict]) -> str:
    if not cands:
        return ""
    primeiro = max(cands, key=lambda c: int(c["votos"]))
    return nome_curto(primeiro["nome"])


def pct(votos: int, validos: int) -> float | None:
    if not validos:
        return None
    return round(100.0 * votos / validos, 2)


def baixar_json(session: requests.Session, url: str, tentativas: int = 5) -> str:
    ultimo: Exception | None = None
    for i in range(tentativas):
        try:
            resp = session.get(url, timeout=60)
            resp.raise_for_status()
            return resp.text
        except (requests.RequestException, ValueError) as exc:
            ultimo = exc
            time.sleep(min(2 ** i, 8))
    raise RuntimeError(f"Falha ao baixar {url}: {ultimo}") from ultimo


def baixar_ano(session: requests.Session, ano: int) -> dict[str, dict]:
    linhas: dict[str, dict] = {}
    for uf in [*UFS, "ZZ", "BR"]:
        url = f"{URL}?{urlencode({'election': ano, 'state': uf})}"
        linhas[uf] = parse_pagina(baixar_json(session, url))
    return linhas


def conferir_totais(bruto: dict[str, dict[str, dict]]) -> None:
    for ano, esperado in TSE_BRASIL.items():
        br = bruto[str(ano)]["BR"]
        fhc = votos_de(br["candidatos"], "fhc")
        lula = votos_de(br["candidatos"], "lula")
        got = {"fhc": fhc, "lula": lula, "vv": int(br["validos"])}
        if got != esperado:
            raise ValueError(f"Totais {ano} divergem do TSE: {got} vs {esperado}.")
        sf = sl = sv = 0
        for uf in [*UFS, "ZZ"]:
            sf += votos_de(bruto[str(ano)][uf]["candidatos"], "fhc")
            sl += votos_de(bruto[str(ano)][uf]["candidatos"], "lula")
            sv += int(bruto[str(ano)][uf]["validos"])
        if (sf, sl, sv) != (fhc, lula, int(br["validos"])):
            raise ValueError(
                f"Soma UFs+exterior {ano} divergente: {(sf, sl, sv)} vs Brasil {(fhc, lula, br['validos'])}."
            )


def carregar_dados(session: requests.Session | None = None, *, forcar: bool = False) -> dict[str, dict[str, dict]]:
    if CACHE.exists() and not forcar:
        bruto = json.loads(CACHE.read_text(encoding="utf-8"))
        conferir_totais(bruto)
        return bruto
    http = session or requests.Session()
    http.headers.update({"User-Agent": UA})
    bruto = {str(ano): baixar_ano(http, ano) for ano in ANOS}
    conferir_totais(bruto)
    CACHE.parent.mkdir(parents=True, exist_ok=True)
    CACHE.write_text(json.dumps(bruto, ensure_ascii=False, indent=2), encoding="utf-8")
    return bruto


def linha_uf(uf: str, bruto: dict[str, dict[str, dict]], cand1: Candidato, cand2: Candidato) -> dict:
    d94 = bruto["1994"][uf]
    d98 = bruto["1998"][uf]
    v94 = int(d94["validos"])
    v98 = int(d98["validos"])
    c1_94 = votos_de(d94["candidatos"], cand1.chave)
    c2_94 = votos_de(d94["candidatos"], cand2.chave)
    c1_98 = votos_de(d98["candidatos"], cand1.chave)
    c2_98 = votos_de(d98["candidatos"], cand2.chave)
    return {
        "uf": uf,
        "unidade": NOMES.get(uf, d94.get("unidade") or uf),
        "regiao": REGIAO.get(uf, ""),
        "y1994_c1_votos": c1_94,
        "y1994_c1_pct": pct(c1_94, v94),
        "y1994_c2_votos": c2_94,
        "y1994_c2_pct": pct(c2_94, v94),
        "y1994_validos": v94,
        "y1994_vencedor": vencedor_local(d94["candidatos"]),
        "y1998_c1_votos": c1_98,
        "y1998_c1_pct": pct(c1_98, v98),
        "y1998_c2_votos": c2_98,
        "y1998_c2_pct": pct(c2_98, v98),
        "y1998_validos": v98,
        "y1998_vencedor": vencedor_local(d98["candidatos"]),
    }


def totais_bloco(ufs: pd.DataFrame, uf: str, unidade: str, regiao: str, cand1: Candidato, cand2: Candidato) -> dict:
    v94 = int(ufs["y1994_validos"].sum())
    v98 = int(ufs["y1998_validos"].sum())
    c1_94 = int(ufs["y1994_c1_votos"].sum())
    c2_94 = int(ufs["y1994_c2_votos"].sum())
    c1_98 = int(ufs["y1998_c1_votos"].sum())
    c2_98 = int(ufs["y1998_c2_votos"].sum())
    return {
        "uf": uf,
        "unidade": unidade,
        "regiao": regiao,
        "y1994_c1_votos": c1_94,
        "y1994_c1_pct": pct(c1_94, v94),
        "y1994_c2_votos": c2_94,
        "y1994_c2_pct": pct(c2_94, v94),
        "y1994_validos": v94,
        "y1994_vencedor": cand1.nome if c1_94 > c2_94 else cand2.nome,
        "y1998_c1_votos": c1_98,
        "y1998_c1_pct": pct(c1_98, v98),
        "y1998_c2_votos": c2_98,
        "y1998_c2_pct": pct(c2_98, v98),
        "y1998_validos": v98,
        "y1998_vencedor": cand1.nome if c1_98 > c2_98 else cand2.nome,
    }


def montar_tabelas(
    bruto: dict[str, dict[str, dict]],
    cand1: Candidato,
    cand2: Candidato,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    linhas = [linha_uf(uf, bruto, cand1, cand2) for uf in UFS]
    ufs = pd.DataFrame(linhas)
    exterior = pd.DataFrame([linha_uf("ZZ", bruto, cand1, cand2)])
    brasil = pd.DataFrame([linha_uf("BR", bruto, cand1, cand2)])
    total = pd.DataFrame(
        [
            totais_bloco(
                ufs,
                "UFs",
                "Total das 27 Unidades da Federação",
                "Brasil (sem exterior)",
                cand1,
                cand2,
            )
        ]
    )
    extra = pd.concat([brasil, total], ignore_index=True)
    return ufs.reset_index(drop=True), exterior.reset_index(drop=True), extra.reset_index(drop=True)


def _borda_fina(ws) -> None:
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = thin


def _titulos(cand1: Candidato, cand2: Candidato) -> list[str]:
    return [
        "UF",
        "Unidade da Federação",
        "Região",
        f"{cand1.nome} (votos válidos)",
        f"{cand1.nome} (% válidos)",
        f"{cand2.nome} (votos válidos)",
        f"{cand2.nome} (% válidos)",
        "Total de votos válidos",
        "Mais votado na UF",
        f"{cand1.nome} (votos válidos)",
        f"{cand1.nome} (% válidos)",
        f"{cand2.nome} (votos válidos)",
        f"{cand2.nome} (% válidos)",
        "Total de votos válidos",
        "Mais votado na UF",
    ]


def _escrever_cabecalho(ws, cand1: Candidato, cand2: Candidato) -> None:
    fill_id = PatternFill("solid", fgColor="1B4F72")
    fill_94 = PatternFill("solid", fgColor="117A65")
    fill_98 = PatternFill("solid", fgColor="6C3483")
    font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("A1:C1")
    ws.merge_cells("D1:I1")
    ws.merge_cells("J1:O1")
    ws["A1"] = "Unidade da Federação"
    ws["D1"] = f"1994 — 1º turno ({DATA[1994]}) — turno único"
    ws["J1"] = f"1998 — 1º turno ({DATA[1998]}) — turno único"
    for col in range(1, 16):
        fill = fill_id if col <= 3 else fill_94 if col <= 9 else fill_98
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = center
    for col, titulo in enumerate(_titulos(cand1, cand2), start=1):
        cell = ws.cell(2, col, titulo)
        fill = fill_id if col <= 3 else fill_94 if col <= 9 else fill_98
        cell.fill = fill
        cell.font = font
        cell.alignment = center
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 36
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:O{ws.max_row}"


def _pintar_vencedor(ws, col_vencedor: int, col_c1: int, col_c2: int, nome1: str, nome2: str) -> None:
    fill_c1 = PatternFill("solid", fgColor="AED6F1")
    fill_c2 = PatternFill("solid", fgColor="F5B7B1")
    fill_outro = PatternFill("solid", fgColor="F9E79F")
    for row in range(3, ws.max_row + 1):
        vencedor = ws.cell(row, col_vencedor).value
        if vencedor == nome1:
            ws.cell(row, col_c1).fill = fill_c1
            ws.cell(row, col_vencedor).fill = fill_c1
        elif vencedor == nome2:
            ws.cell(row, col_c2).fill = fill_c2
            ws.cell(row, col_vencedor).fill = fill_c2
        elif vencedor:
            ws.cell(row, col_vencedor).fill = fill_outro


def _formatar_corpo(ws) -> None:
    inteiros = {4, 6, 8, 10, 12, 14}
    percentuais = {5, 7, 11, 13}
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    for row in range(3, ws.max_row + 1):
        for col in range(1, 16):
            ws.cell(row, col).alignment = center
        ws.cell(row, 2).alignment = left
        for col in inteiros:
            ws.cell(row, col).number_format = "#,##0"
        for col in percentuais:
            ws.cell(row, col).number_format = "0.00"
    larguras = [8, 38, 22, 22, 16, 18, 16, 22, 18, 22, 16, 18, 16, 22, 18]
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _preencher_aba(
    ws,
    df: pd.DataFrame,
    cand1: Candidato,
    cand2: Candidato,
    destacar_ultima: bool = False,
) -> None:
    corpo = df.loc[:, COLUNAS]
    for r_idx, valores in enumerate(corpo.itertuples(index=False, name=None), start=3):
        for c_idx, value in enumerate(valores, start=1):
            ws.cell(r_idx, c_idx, value)
    _escrever_cabecalho(ws, cand1, cand2)
    _formatar_corpo(ws)
    _pintar_vencedor(ws, 9, 4, 6, cand1.nome, cand2.nome)
    _pintar_vencedor(ws, 15, 10, 12, cand1.nome, cand2.nome)
    if destacar_ultima and ws.max_row >= 3:
        fill = PatternFill("solid", fgColor="D5F5E3")
        font = Font(bold=True)
        for col in range(1, 16):
            cell = ws.cell(ws.max_row, col)
            cell.fill = fill
            cell.font = font
    _borda_fina(ws)


def fmt_int_br(valor: int) -> str:
    return f"{int(valor):,}".replace(",", ".")


def _notas(cand1: Candidato, cand2: Candidato) -> list[str]:
    t94 = TSE_BRASIL[1994]
    t98 = TSE_BRASIL[1998]
    return [
        "Tribunal Superior Eleitoral (TSE) — 1º turno das eleições presidenciais de 1994 e 1998 (turno único em ambos os pleitos).",
        (
            "Os dois candidatos mais votados no 1º turno nacional, nos dois anos: "
            f"{cand1.nome} (Fernando Henrique Cardoso, PSDB, nº 45) e {cand2.nome} "
            "(Luiz Inácio Lula da Silva, PT, nº 13)."
        ),
        "A planilha reporta os votos válidos desses dois candidatos em cada uma das 27 Unidades da Federação.",
        "Em 1998, Ciro Gomes (PPS) foi o mais votado no Ceará; a coluna “Mais votado na UF” registra o primeiro colocado local, que pode diferir dos dois primeiros nacionais.",
        "Votos válidos: votos nominais (excluídos brancos e nulos). Percentuais sobre o total de válidos da UF no respectivo ano.",
        "Aba Por_UF: 26 estados + Distrito Federal. A última linha soma só as 27 UFs (sem voto no exterior).",
        "Aba Brasil_e_exterior: total nacional (inclui ZZ/exterior), voto no exterior e a soma das 27 UFs.",
        (
            "Microdados por UF: http://electionresources.org/br/president.php (tabelas compiladas da totalização do TSE). "
            "Páginas TSE 1994: https://www.tse.jus.br/eleicoes/eleicoes-anteriores/eleicoes-1994/resultados-das-eleicoes-1994"
        ),
        (
            f"Brasil 1994: {cand1.nome} {fmt_int_br(t94['fhc'])} (54,28%); "
            f"{cand2.nome} {fmt_int_br(t94['lula'])} (27,04%); válidos {fmt_int_br(t94['vv'])}."
        ),
        (
            f"Brasil 1998: {cand1.nome} {fmt_int_br(t98['fhc'])} (53,06%); "
            f"{cand2.nome} {fmt_int_br(t98['lula'])} (31,71%); válidos {fmt_int_br(t98['vv'])}."
        ),
        "A soma das 27 UFs + voto no exterior reproduz o total Brasil nos dois anos.",
        "O TSE alerta que os arquivos brutos de 1994–1998 podem estar incompletos nas bases centralizadas; as tabelas de totalização por UF aqui usadas conferem com as páginas oficiais de resultado (ex.: SP, RS, DF, SC e CE em 1994).",
    ]


def gravar_xlsx(
    ufs: pd.DataFrame,
    exterior: pd.DataFrame,
    extra: pd.DataFrame,
    cand1: Candidato,
    cand2: Candidato,
    caminho: Path,
) -> Path:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Por_UF"
    ufs_com_total = pd.concat([ufs, extra.loc[extra["uf"] == "UFs"]], ignore_index=True)
    _preencher_aba(ws, ufs_com_total, cand1, cand2, destacar_ultima=True)
    ws.sheet_view.showGridLines = False

    ws2 = wb.create_sheet("Brasil_e_exterior")
    brasil_ext = pd.concat(
        [extra.loc[extra["uf"] == "BR"], exterior, extra.loc[extra["uf"] == "UFs"]],
        ignore_index=True,
    )
    _preencher_aba(ws2, brasil_ext, cand1, cand2)
    ws2.sheet_view.showGridLines = False

    ws3 = wb.create_sheet("Fonte")
    ws3["A1"] = "Fonte, recorte e conceito"
    ws3["A1"].font = Font(bold=True, size=14)
    for i, texto in enumerate(_notas(cand1, cand2), start=3):
        ws3[f"A{i}"] = texto
        ws3[f"A{i}"].alignment = Alignment(wrap_text=True)
        ws3.row_dimensions[i].height = 18
    ws3.column_dimensions["A"].width = 140
    wb.save(caminho)
    return caminho


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Planilha: votos válidos dos dois candidatos à Presidência mais votados "
            "no 1º turno de 1994 e 1998, em cada UF."
        )
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "output" / "eleicoes_1994_1998_presidente_1t_por_uf.xlsx",
    )
    parser.add_argument("--baixar", action="store_true", help="Força novo download.")
    args = parser.parse_args(argv)
    bruto = carregar_dados(forcar=args.baixar)
    cand1, cand2 = dois_mais_votados(bruto["1994"]["BR"])
    c1_98, c2_98 = dois_mais_votados(bruto["1998"]["BR"])
    if (c1_98.chave, c2_98.chave) != (cand1.chave, cand2.chave):
        raise ValueError(
            f"Top 2 de 1998 ({c1_98.nome}, {c2_98.nome}) diferem de 1994 ({cand1.nome}, {cand2.nome})."
        )
    ufs, exterior, extra = montar_tabelas(bruto, cand1, cand2)
    caminho = gravar_xlsx(ufs, exterior, extra, cand1, cand2, args.output)
    print(f"[OK] {caminho}")
    print(f"Candidatos: {cand1.nome} e {cand2.nome} (1º turno 1994 e 1998)")
    print(ufs.to_string(index=False))
    br = extra.loc[extra["uf"] == "BR"].iloc[0]
    print("Brasil 1994:", int(br["y1994_c1_votos"]), int(br["y1994_c2_votos"]), "válidos", int(br["y1994_validos"]))
    print("Brasil 1998:", int(br["y1998_c1_votos"]), int(br["y1998_c2_votos"]), "válidos", int(br["y1998_validos"]))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
