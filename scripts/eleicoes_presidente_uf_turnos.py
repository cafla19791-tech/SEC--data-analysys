#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Votos válidos dos dois candidatos à Presidência mais votados, por UF.

Consolida 1º e 2º turnos das Eleições Gerais 2022 (cargo de presidente) em
cada uma das 27 Unidades da Federação, a partir da totalização oficial do TSE.

Fonte:
  https://resultados.tse.jus.br/oficial/ele2022/{eleicao}/dados-simplificados/

Uso::

  PYTHONPATH=. python3 scripts/eleicoes_presidente_uf_turnos.py
  PYTHONPATH=. python3 scripts/eleicoes_presidente_uf_turnos.py --output output/eleicoes.xlsx
"""

from __future__ import annotations

import argparse
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
import requests
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

CARGO = "0001"
ELEICAO_1T = 544
ELEICAO_2T = 545
DATA_1T = "02/10/2022"
DATA_2T = "30/10/2022"
UA = "SEC-data-analysys/eleicoes-presidente-uf-turnos"

# Totais nacionais publicados pelo TSE após 100% das seções.
TSE_BRASIL_1T = {
    "13": 57_259_504,
    "22": 51_072_345,
    "vv": 118_229_719,
}
TSE_BRASIL_2T = {
    "13": 60_345_999,
    "22": 58_206_354,
    "vv": 118_552_353,
}

UFS = [
    "ac", "al", "ap", "am", "ba", "ce", "df", "es", "go", "ma", "mt", "ms",
    "mg", "pa", "pb", "pr", "pe", "pi", "rj", "rn", "rs", "ro", "rr", "sc",
    "sp", "se", "to",
]

NOMES = {
    "AC": "Acre",
    "AL": "Alagoas",
    "AP": "Amapá",
    "AM": "Amazonas",
    "BA": "Bahia",
    "CE": "Ceará",
    "DF": "Distrito Federal",
    "ES": "Espírito Santo",
    "GO": "Goiás",
    "MA": "Maranhão",
    "MT": "Mato Grosso",
    "MS": "Mato Grosso do Sul",
    "MG": "Minas Gerais",
    "PA": "Pará",
    "PB": "Paraíba",
    "PR": "Paraná",
    "PE": "Pernambuco",
    "PI": "Piauí",
    "RJ": "Rio de Janeiro",
    "RN": "Rio Grande do Norte",
    "RS": "Rio Grande do Sul",
    "RO": "Rondônia",
    "RR": "Roraima",
    "SC": "Santa Catarina",
    "SP": "São Paulo",
    "SE": "Sergipe",
    "TO": "Tocantins",
    "ZZ": "Exterior",
    "BR": "Brasil",
}

REGIAO = {
    "AC": "Norte",
    "AL": "Nordeste",
    "AP": "Norte",
    "AM": "Norte",
    "BA": "Nordeste",
    "CE": "Nordeste",
    "DF": "Centro-Oeste",
    "ES": "Sudeste",
    "GO": "Centro-Oeste",
    "MA": "Nordeste",
    "MT": "Centro-Oeste",
    "MS": "Centro-Oeste",
    "MG": "Sudeste",
    "PA": "Norte",
    "PB": "Nordeste",
    "PR": "Sul",
    "PE": "Nordeste",
    "PI": "Nordeste",
    "RJ": "Sudeste",
    "RN": "Nordeste",
    "RS": "Sul",
    "RO": "Norte",
    "RR": "Norte",
    "SC": "Sul",
    "SP": "Sudeste",
    "SE": "Nordeste",
    "TO": "Norte",
    "ZZ": "Exterior",
    "BR": "Brasil",
}

NOME_EXIBICAO = {
    "13": "Lula",
    "22": "Jair Bolsonaro",
}

URL_JSON = (
    "https://resultados.tse.jus.br/oficial/ele2022/{eleicao}/"
    "dados-simplificados/{uf}/{uf}-c{cargo}-e{eleicao:06d}-r.json"
)

COLUNAS_UF = [
    "uf",
    "unidade",
    "regiao",
    "t1_c1_votos",
    "t1_c1_pct",
    "t1_c2_votos",
    "t1_c2_pct",
    "t1_votos_validos",
    "t1_vencedor",
    "t2_c1_votos",
    "t2_c1_pct",
    "t2_c2_votos",
    "t2_c2_pct",
    "t2_votos_validos",
    "t2_vencedor",
]


@dataclass(frozen=True)
class Candidato:
    numero: str
    nome_tse: str
    nome: str
    coligacao: str
    votos_brasil_1t: int
    pct_brasil_1t: float


def pct_oficial(texto: str) -> float:
    """Percentual TSE (vírgula decimal, ponto de milhar)."""
    return float(str(texto).replace(".", "").replace(",", "."))


def votos_int(valor: Any) -> int:
    return int(str(valor).replace(".", "").replace(",", ""))


def fmt_int_br(valor: int) -> str:
    return f"{int(valor):,}".replace(",", ".")


def fmt_pct_br(valor: float) -> str:
    return f"{valor:.2f}".replace(".", ",") + "%"


def nome_candidato(numero: str, nome_tse: str) -> str:
    return NOME_EXIBICAO.get(str(numero), str(nome_tse).title())


def url_turno(eleicao: int, uf: str) -> str:
    return URL_JSON.format(eleicao=eleicao, uf=uf.lower(), cargo=CARGO)


def dois_mais_votados(payload: dict) -> tuple[Candidato, Candidato]:
    """Os dois candidatos com mais votos válidos no 1º turno nacional."""
    candidatos = sorted(payload["cand"], key=lambda c: votos_int(c["vap"]), reverse=True)
    if len(candidatos) < 2:
        raise ValueError("JSON do TSE sem dois candidatos à Presidência.")
    top: list[Candidato] = []
    for item in candidatos[:2]:
        numero = str(item["n"])
        top.append(
            Candidato(
                numero=numero,
                nome_tse=str(item["nm"]),
                nome=nome_candidato(numero, str(item["nm"])),
                coligacao=str(item.get("cc") or ""),
                votos_brasil_1t=votos_int(item["vap"]),
                pct_brasil_1t=pct_oficial(item["pvap"]),
            )
        )
    return top[0], top[1]


def candidato_do_payload(payload: dict, numero: str) -> dict:
    for item in payload["cand"]:
        if str(item["n"]) == str(numero):
            return item
    raise KeyError(f"Candidato nº {numero} ausente em {payload.get('cdabr')}.")


def vencedor_por_votos(v1: int, v2: int, nome1: str, nome2: str) -> str:
    if v1 > v2:
        return nome1
    if v2 > v1:
        return nome2
    return "Empate"


def parse_abrangencia(payload: dict, cand1: Candidato, cand2: Candidato) -> dict:
    c1 = candidato_do_payload(payload, cand1.numero)
    c2 = candidato_do_payload(payload, cand2.numero)
    v1 = votos_int(c1["vap"])
    v2 = votos_int(c2["vap"])
    uf = str(payload["cdabr"]).upper()
    return {
        "uf": uf,
        "unidade": NOMES.get(uf, uf),
        "regiao": REGIAO.get(uf, ""),
        "c1_votos": v1,
        "c1_pct": pct_oficial(c1["pvap"]),
        "c2_votos": v2,
        "c2_pct": pct_oficial(c2["pvap"]),
        "votos_validos": votos_int(payload["vv"]),
        "votos_brancos": votos_int(payload.get("vb") or 0),
        "votos_nulos": votos_int(payload.get("vn") or 0),
        "comparecimento": votos_int(payload.get("c") or 0),
        "abstencao": votos_int(payload.get("a") or 0),
        "eleitores_aptos": votos_int(payload.get("e") or 0),
        "vencedor": vencedor_por_votos(v1, v2, cand1.nome, cand2.nome),
        "secoes_totalizadas_pct": pct_oficial(payload.get("pst") or "0"),
    }


def fundir_turnos(t1: dict, t2: dict) -> dict:
    if t1["uf"] != t2["uf"]:
        raise ValueError(f"UF divergente: {t1['uf']} vs {t2['uf']}")
    return {
        "uf": t1["uf"],
        "unidade": t1["unidade"],
        "regiao": t1["regiao"],
        "t1_c1_votos": t1["c1_votos"],
        "t1_c1_pct": t1["c1_pct"],
        "t1_c2_votos": t1["c2_votos"],
        "t1_c2_pct": t1["c2_pct"],
        "t1_votos_validos": t1["votos_validos"],
        "t1_vencedor": t1["vencedor"],
        "t2_c1_votos": t2["c1_votos"],
        "t2_c1_pct": t2["c1_pct"],
        "t2_c2_votos": t2["c2_votos"],
        "t2_c2_pct": t2["c2_pct"],
        "t2_votos_validos": t2["votos_validos"],
        "t2_vencedor": t2["vencedor"],
        "t1_brancos": t1["votos_brancos"],
        "t1_nulos": t1["votos_nulos"],
        "t1_comparecimento": t1["comparecimento"],
        "t1_abstencao": t1["abstencao"],
        "t1_aptos": t1["eleitores_aptos"],
        "t2_brancos": t2["votos_brancos"],
        "t2_nulos": t2["votos_nulos"],
        "t2_comparecimento": t2["comparecimento"],
        "t2_abstencao": t2["abstencao"],
        "t2_aptos": t2["eleitores_aptos"],
        "t1_secoes_pct": t1["secoes_totalizadas_pct"],
        "t2_secoes_pct": t2["secoes_totalizadas_pct"],
    }


def baixar_json(session: requests.Session, url: str, tentativas: int = 5) -> dict:
    ultimo: Exception | None = None
    for i in range(tentativas):
        try:
            resp = session.get(url, timeout=60)
            resp.raise_for_status()
            return resp.json()
        except (requests.RequestException, ValueError) as exc:
            ultimo = exc
            time.sleep(min(2 ** i, 8))
    raise RuntimeError(f"Falha ao baixar {url}: {ultimo}") from ultimo


def nova_sessao() -> requests.Session:
    session = requests.Session()
    session.headers.update({"User-Agent": UA, "Accept": "application/json"})
    return session


def conferir_totais_nacionais(
    cand1: Candidato,
    cand2: Candidato,
    br_1t: dict,
    br_2t: dict,
) -> None:
    """Garante que o JSON bate com a totalização publicada pelo TSE."""
    got_1t = {
        cand1.numero: votos_int(candidato_do_payload(br_1t, cand1.numero)["vap"]),
        cand2.numero: votos_int(candidato_do_payload(br_1t, cand2.numero)["vap"]),
        "vv": votos_int(br_1t["vv"]),
    }
    got_2t = {
        cand1.numero: votos_int(candidato_do_payload(br_2t, cand1.numero)["vap"]),
        cand2.numero: votos_int(candidato_do_payload(br_2t, cand2.numero)["vap"]),
        "vv": votos_int(br_2t["vv"]),
    }
    if got_1t != TSE_BRASIL_1T or got_2t != TSE_BRASIL_2T:
        raise ValueError(
            "Totais nacionais do JSON divergem da totalização oficial do TSE "
            f"(1º turno {got_1t} vs {TSE_BRASIL_1T}; 2º turno {got_2t} vs {TSE_BRASIL_2T})."
        )


def baixar_tse(
    session: requests.Session | None = None,
) -> tuple[Candidato, Candidato, list[dict], dict]:
    http = session or nova_sessao()
    br_1t = baixar_json(http, url_turno(ELEICAO_1T, "br"))
    br_2t = baixar_json(http, url_turno(ELEICAO_2T, "br"))
    cand1, cand2 = dois_mais_votados(br_1t)
    conferir_totais_nacionais(cand1, cand2, br_1t, br_2t)

    linhas: list[dict] = []
    for uf in [*UFS, "zz"]:
        p1 = baixar_json(http, url_turno(ELEICAO_1T, uf))
        p2 = baixar_json(http, url_turno(ELEICAO_2T, uf))
        linhas.append(
            fundir_turnos(
                parse_abrangencia(p1, cand1, cand2),
                parse_abrangencia(p2, cand1, cand2),
            )
        )
    brasil = fundir_turnos(
        parse_abrangencia(br_1t, cand1, cand2),
        parse_abrangencia(br_2t, cand1, cand2),
    )
    return cand1, cand2, linhas, brasil


def _totais_bloco(
    origem: pd.DataFrame,
    uf: str,
    unidade: str,
    regiao: str,
    cand1: Candidato,
    cand2: Candidato,
) -> dict:
    t1v = int(origem["t1_votos_validos"].sum())
    t2v = int(origem["t2_votos_validos"].sum())
    t1c1 = int(origem["t1_c1_votos"].sum())
    t1c2 = int(origem["t1_c2_votos"].sum())
    t2c1 = int(origem["t2_c1_votos"].sum())
    t2c2 = int(origem["t2_c2_votos"].sum())
    return {
        "uf": uf,
        "unidade": unidade,
        "regiao": regiao,
        "t1_c1_votos": t1c1,
        "t1_c1_pct": round(100.0 * t1c1 / t1v, 2) if t1v else None,
        "t1_c2_votos": t1c2,
        "t1_c2_pct": round(100.0 * t1c2 / t1v, 2) if t1v else None,
        "t1_votos_validos": t1v,
        "t1_vencedor": vencedor_por_votos(t1c1, t1c2, cand1.nome, cand2.nome),
        "t2_c1_votos": t2c1,
        "t2_c1_pct": round(100.0 * t2c1 / t2v, 2) if t2v else None,
        "t2_c2_votos": t2c2,
        "t2_c2_pct": round(100.0 * t2c2 / t2v, 2) if t2v else None,
        "t2_votos_validos": t2v,
        "t2_vencedor": vencedor_por_votos(t2c1, t2c2, cand1.nome, cand2.nome),
    }


def conferir_soma_ufs_mais_exterior(ufs: pd.DataFrame, exterior: pd.DataFrame, brasil: dict) -> None:
    """A soma das 27 UFs + ZZ deve reproduzir o total nacional do TSE."""
    campos = [
        "t1_c1_votos",
        "t1_c2_votos",
        "t1_votos_validos",
        "t2_c1_votos",
        "t2_c2_votos",
        "t2_votos_validos",
    ]
    for campo in campos:
        soma = int(ufs[campo].sum()) + int(exterior[campo].sum())
        oficial = int(brasil[campo])
        if soma != oficial:
            raise ValueError(
                f"Soma UFs+exterior divergente em {campo}: {soma} vs Brasil {oficial}."
            )


def montar_tabelas(
    linhas: list[dict],
    brasil: dict,
    cand1: Candidato,
    cand2: Candidato,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    todos = pd.DataFrame(linhas)
    ufs = todos.loc[todos["uf"] != "ZZ", COLUNAS_UF].copy()
    exterior = todos.loc[todos["uf"] == "ZZ", COLUNAS_UF].copy()
    total_ufs = _totais_bloco(
        ufs,
        "UFs",
        "Total das 27 Unidades da Federação",
        "Brasil (sem exterior)",
        cand1,
        cand2,
    )
    brasil_row = {col: brasil[col] for col in COLUNAS_UF}
    conferir_soma_ufs_mais_exterior(ufs, exterior, brasil_row)
    extra = pd.DataFrame([brasil_row, total_ufs])
    return ufs.reset_index(drop=True), exterior.reset_index(drop=True), extra.reset_index(drop=True)


def _borda_fina(ws) -> None:
    thin = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = thin


def _titulos(cand1: Candidato, cand2: Candidato) -> list[str]:
    return [
        "UF",
        "Unidade da Federação",
        "Região",
        f"{cand1.nome} (votos válidos)",
        f"{cand1.nome} (% válidos)",
        f"{cand2.nome} (votos válidos)",
        f"{cand2.nome} (% válidos)",
        "Total de votos válidos",
        "Mais votado",
        f"{cand1.nome} (votos válidos)",
        f"{cand1.nome} (% válidos)",
        f"{cand2.nome} (votos válidos)",
        f"{cand2.nome} (% válidos)",
        "Total de votos válidos",
        "Mais votado",
    ]


def _escrever_cabecalho_duplo(ws, cand1: Candidato, cand2: Candidato) -> None:
    fill_id = PatternFill("solid", fgColor="1B4F72")
    fill_t1 = PatternFill("solid", fgColor="117A65")
    fill_t2 = PatternFill("solid", fgColor="6C3483")
    font = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells("A1:C1")
    ws.merge_cells("D1:I1")
    ws.merge_cells("J1:O1")
    ws["A1"] = "Unidade da Federação"
    ws["D1"] = f"1º turno — {DATA_1T}"
    ws["J1"] = f"2º turno — {DATA_2T}"
    for col in range(1, 16):
        if col <= 3:
            fill = fill_id
        elif col <= 9:
            fill = fill_t1
        else:
            fill = fill_t2
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = center
    titulos = _titulos(cand1, cand2)
    for col, titulo in enumerate(titulos, start=1):
        cell = ws.cell(2, col, titulo)
        if col <= 3:
            fill = fill_id
        elif col <= 9:
            fill = fill_t1
        else:
            fill = fill_t2
        cell.fill = fill
        cell.font = font
        cell.alignment = center
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 36
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:O{ws.max_row}"


def _pintar_vencedor(ws, col_vencedor: int, col_c1: int, col_c2: int, nome1: str, nome2: str) -> None:
    fill_c1 = PatternFill("solid", fgColor="F5B7B1")
    fill_c2 = PatternFill("solid", fgColor="AED6F1")
    for row in range(3, ws.max_row + 1):
        vencedor = ws.cell(row, col_vencedor).value
        if vencedor == nome1:
            ws.cell(row, col_c1).fill = fill_c1
            ws.cell(row, col_vencedor).fill = fill_c1
        elif vencedor == nome2:
            ws.cell(row, col_c2).fill = fill_c2
            ws.cell(row, col_vencedor).fill = fill_c2


def _formatar_corpo(ws, primeira: int = 3) -> None:
    inteiros = {4, 6, 8, 10, 12, 14}
    percentuais = {5, 7, 11, 13}
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    for row in range(primeira, ws.max_row + 1):
        for col in range(1, 16):
            ws.cell(row, col).alignment = center
        ws.cell(row, 2).alignment = left
        for col in inteiros:
            ws.cell(row, col).number_format = "#,##0"
        for col in percentuais:
            ws.cell(row, col).number_format = "0.00"
    larguras = [8, 38, 22, 22, 16, 24, 16, 22, 16, 22, 16, 24, 16, 22, 16]
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _preencher_aba(
    ws,
    df: pd.DataFrame,
    cand1: Candidato,
    cand2: Candidato,
    destacar_ultima: bool = False,
) -> None:
    corpo = df.loc[:, COLUNAS_UF]
    for r_idx, valores in enumerate(corpo.itertuples(index=False, name=None), start=3):
        for c_idx, value in enumerate(valores, start=1):
            ws.cell(r_idx, c_idx, value)
    _escrever_cabecalho_duplo(ws, cand1, cand2)
    _formatar_corpo(ws)
    _pintar_vencedor(ws, 9, 4, 6, cand1.nome, cand2.nome)
    _pintar_vencedor(ws, 15, 10, 12, cand1.nome, cand2.nome)
    if destacar_ultima and ws.max_row >= 3:
        fill = PatternFill("solid", fgColor="D5F5E3")
        font = Font(bold=True)
        for col in range(1, 16):
            cell = ws.cell(ws.max_row, col)
            cell.fill = fill
            cell.font = font
    _borda_fina(ws)


def _notas_fonte(cand1: Candidato, cand2: Candidato) -> list[str]:
    pct2_c1 = 100.0 * TSE_BRASIL_2T[cand1.numero] / TSE_BRASIL_2T["vv"]
    pct2_c2 = 100.0 * TSE_BRASIL_2T[cand2.numero] / TSE_BRASIL_2T["vv"]
    pct1_c2 = 100.0 * TSE_BRASIL_1T[cand2.numero] / TSE_BRASIL_1T["vv"]
    return [
        "Tribunal Superior Eleitoral (TSE) — totalização oficial das Eleições Gerais 2022, cargo de Presidente da República.",
        (
            "Os dois candidatos mais votados no 1º turno nacional (e únicos no 2º turno): "
            f"{cand1.nome} (nº {cand1.numero}) e {cand2.nome} (nº {cand2.numero})."
        ),
        f"{cand1.nome}: {cand1.nome_tse}. Coligação/partido: {cand1.coligacao}.",
        f"{cand2.nome}: {cand2.nome_tse}. Coligação/partido: {cand2.coligacao}.",
        f"1º turno: {DATA_1T} (eleição {ELEICAO_1T}). 2º turno: {DATA_2T} (eleição {ELEICAO_2T}).",
        (
            "Votos válidos: votos nominais nos candidatos (excluídos brancos e nulos). "
            "Os percentuais são sobre o total de válidos da UF no respectivo turno."
        ),
        (
            "Aba Por_UF: as 26 Unidades da Federação + o Distrito Federal (27 UFs). "
            "A última linha soma só essas 27 (sem voto no exterior)."
        ),
        (
            "Aba Brasil_e_exterior: total nacional oficial do TSE (inclui ZZ/exterior), "
            "voto no exterior e a soma das 27 UFs."
        ),
        (
            "Arquivos: https://resultados.tse.jus.br/oficial/ele2022/{544|545}/"
            "dados-simplificados/{uf}/{uf}-c0001-e000{544|545}-r.json"
        ),
        (
            f"Brasil oficial 1º turno: {cand1.nome} {fmt_int_br(cand1.votos_brasil_1t)} "
            f"({fmt_pct_br(cand1.pct_brasil_1t)}); "
            f"{cand2.nome} {fmt_int_br(TSE_BRASIL_1T[cand2.numero])} ({fmt_pct_br(pct1_c2)}); "
            f"válidos {fmt_int_br(TSE_BRASIL_1T['vv'])}."
        ),
        (
            f"Brasil oficial 2º turno: {cand1.nome} {fmt_int_br(TSE_BRASIL_2T[cand1.numero])} "
            f"({fmt_pct_br(pct2_c1)}); "
            f"{cand2.nome} {fmt_int_br(TSE_BRASIL_2T[cand2.numero])} ({fmt_pct_br(pct2_c2)}); "
            f"válidos {fmt_int_br(TSE_BRASIL_2T['vv'])}."
        ),
        "Notícias TSE da totalização (100% das seções): 1º turno em 04/10/2022 e 2º turno em 31/10/2022.",
    ]


def gravar_xlsx(
    ufs: pd.DataFrame,
    exterior: pd.DataFrame,
    extra: pd.DataFrame,
    cand1: Candidato,
    cand2: Candidato,
    caminho: Path,
) -> Path:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Por_UF"
    ufs_com_total = pd.concat([ufs, extra.loc[extra["uf"] == "UFs"]], ignore_index=True)
    _preencher_aba(ws, ufs_com_total, cand1, cand2, destacar_ultima=True)
    ws.sheet_view.showGridLines = False

    ws2 = wb.create_sheet("Brasil_e_exterior")
    brasil_ext = pd.concat(
        [extra.loc[extra["uf"] == "BR"], exterior, extra.loc[extra["uf"] == "UFs"]],
        ignore_index=True,
    )
    _preencher_aba(ws2, brasil_ext, cand1, cand2)
    ws2.sheet_view.showGridLines = False

    ws3 = wb.create_sheet("Fonte")
    ws3["A1"] = "Fonte, recorte e conceito"
    ws3["A1"].font = Font(bold=True, size=14)
    notas = _notas_fonte(cand1, cand2)
    for i, texto in enumerate(notas, start=3):
        ws3[f"A{i}"] = texto
        ws3[f"A{i}"].alignment = Alignment(wrap_text=True)
    ws3.column_dimensions["A"].width = 140
    ws3.row_dimensions[1].height = 22
    for row in range(3, 3 + len(notas)):
        ws3.row_dimensions[row].height = 18

    wb.save(caminho)
    return caminho


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Planilha TSE: votos válidos dos dois candidatos à Presidência "
            "mais votados, 1º e 2º turnos, em cada UF."
        )
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "output" / "eleicoes_2022_presidente_1t_2t_por_uf.xlsx",
    )
    args = parser.parse_args(argv)
    cand1, cand2, linhas, brasil = baixar_tse()
    ufs, exterior, extra = montar_tabelas(linhas, brasil, cand1, cand2)
    caminho = gravar_xlsx(ufs, exterior, extra, cand1, cand2, args.output)
    print(f"[OK] {caminho}")
    print(f"Candidatos: {cand1.nome} (nº {cand1.numero}) e {cand2.nome} (nº {cand2.numero})")
    print(ufs.to_string(index=False))
    br = extra.loc[extra["uf"] == "BR"].iloc[0]
    print(
        "Brasil TSE 1º turno:",
        int(br["t1_c1_votos"]),
        int(br["t1_c2_votos"]),
        "válidos",
        int(br["t1_votos_validos"]),
    )
    print(
        "Brasil TSE 2º turno:",
        int(br["t2_c1_votos"]),
        int(br["t2_c2_votos"]),
        "válidos",
        int(br["t2_votos_validos"]),
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
