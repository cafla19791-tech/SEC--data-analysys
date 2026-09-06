#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fatores condicionantes da Dívida Bruta do Governo Geral, 1995–2026.

O Banco Central publica a decomposição oficial da DBGG (juros nominais,
emissões líquidas, câmbio, reconhecimento/privatizações e efeito PIB) nas
Notas de Estatísticas Fiscais, mas esses fluxos **não** são séries SGS.

Este script reconstitui a identidade em p.p. do PIB a partir das séries
públicas:

* estoque DBGG — SGS 13762/13761 (metodologia a partir de 2008; desde
  dez/2006) e SGS 4537 (metodologia até 2007; desde 2002);
* juros e primário do Governo Geral — NFSP sem desvalorização cambial
  (5751+5753 e 5784+5786 em 12 meses; 4607+4610+4611 e 4640+4643+4644
  no mensal, para anos incompletos);
* efeito do PIB nominal — −d_{t−1} × g/(1+g), com g pela série 4382
  (PIB acumulado em 12 meses) no último mês de cada ano.

O resíduo («Demais fatores») fecha a identidade e absorve câmbio,
reconhecimento de dívidas, privatizações, a diferença entre emissões
líquidas da DBGG e o primário da NFSP, e a diferença entre juros da
dívida bruta e juros da dívida líquida.

Sinal: valor positivo aumenta a relação DBGG/PIB (necessidade primária
positiva = déficit). Superávit primário entra com sinal negativo.

Uso::

  python scripts/fatores_condicionantes_dbgg.py
  python fatores_condicionantes_dbgg.py --saida output/fatores.xlsx
"""

from __future__ import annotations

import argparse
import sys
import time
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

MARKER = "fatores-dbgg-1995-2026-20260906"
BCB_SGS = "https://api.bcb.gov.br/dados/serie/bcdata.sgs.{cod}/dados"
ANO_INICIO = 1995
ANO_FIM = 2026
# Resíduo anual em p.p. do PIB (arredondamento das séries publicadas).
TOLERANCIA_PP = 0.05

COL_JUROS = "Juros nominais"
COL_PRIMARIO = "Primário"
COL_PIB = "Efeito do PIB nominal"
COL_DEMAIS = "Demais fatores"
COL_VARIACAO = "Variação da DBGG/PIB"
COL_ESTOQUE_PP = "Estoque DBGG (% PIB)"
COL_ESTOQUE_RS = "Estoque DBGG (R$ milhões)"
COL_METODO = "Metodologia"

FATORES = (COL_JUROS, COL_PRIMARIO, COL_PIB, COL_DEMAIS)

METODO_ATUAL = "a partir de 2008"
METODO_ANTIGO = "até 2007"


@dataclass(frozen=True)
class Serie:
    codigo: int
    nome: str
    papel: str  # estoque_pp | estoque_rs | pib | nfsp_12m | nfsp_mes


SERIES: tuple[Serie, ...] = (
    Serie(13762, "DBGG (% PIB) — metodologia a partir de 2008", "estoque_pp"),
    Serie(13761, "DBGG (R$ milhões) — metodologia a partir de 2008", "estoque_rs"),
    Serie(4537, "DBGG (% PIB) — metodologia até 2007", "estoque_pp"),
    Serie(4382, "PIB acumulado em 12 meses (R$ milhões)", "pib"),
    Serie(5751, "NFSP juros 12 meses — Governo Federal (% PIB)", "nfsp_12m"),
    Serie(5753, "NFSP juros 12 meses — Estados e municípios (% PIB)", "nfsp_12m"),
    Serie(5784, "NFSP primário 12 meses — Governo Federal (% PIB)", "nfsp_12m"),
    Serie(5786, "NFSP primário 12 meses — Estados e municípios (% PIB)", "nfsp_12m"),
    Serie(4607, "NFSP juros mensal — Governo Federal (R$ milhões)", "nfsp_mes"),
    Serie(4610, "NFSP juros mensal — Governos estaduais (R$ milhões)", "nfsp_mes"),
    Serie(4611, "NFSP juros mensal — Governos municipais (R$ milhões)", "nfsp_mes"),
    Serie(4640, "NFSP primário mensal — Governo Federal (R$ milhões)", "nfsp_mes"),
    Serie(4643, "NFSP primário mensal — Governos estaduais (R$ milhões)", "nfsp_mes"),
    Serie(4644, "NFSP primário mensal — Governos municipais (R$ milhões)", "nfsp_mes"),
)

CODIGOS = tuple(s.codigo for s in SERIES)
# Início mínimo aceitável do cache (abaixo disso, baixa de novo).
CACHE_MINIMO: dict[int, date] = {
    13762: date(2006, 12, 1),
    13761: date(2006, 12, 1),
    4537: date(2001, 12, 1),
    4382: date(1991, 1, 1),
    5751: date(1999, 12, 1),
    5753: date(1991, 12, 1),
    5784: date(1999, 12, 1),
    5786: date(1991, 12, 1),
    4607: date(1999, 1, 1),
    4610: date(1998, 1, 1),
    4611: date(1998, 1, 1),
    4640: date(1999, 1, 1),
    4643: date(1998, 1, 1),
    4644: date(1998, 1, 1),
}


def _http_get(url: str, params: dict, tentativas: int = 5) -> list:
    ultimo: Exception | None = None
    headers = {"User-Agent": f"SEC-data-analysys/{MARKER}"}
    for i in range(tentativas):
        try:
            resp = requests.get(url, params=params, headers=headers, timeout=120)
            if resp.status_code == 404:
                return []
            if resp.status_code != 200 or not resp.text.strip():
                raise RuntimeError(f"HTTP {resp.status_code} vazio")
            if resp.text.lstrip().startswith("<") or resp.text.lstrip().startswith("<?xml"):
                raise RuntimeError("resposta XML/HTML em vez de JSON")
            dados = resp.json()
            if not isinstance(dados, list):
                raise RuntimeError(f"JSON inesperado: {type(dados)}")
            return dados
        except Exception as exc:  # noqa: BLE001 — retry de rede/SGS
            ultimo = exc
            time.sleep(1.5 * (2**i))
    raise RuntimeError(f"falha ao baixar {url} {params}: {ultimo}") from ultimo


def baixar_sgs(
    codigo: int,
    inicio: date,
    fim: date,
    *,
    cache: Path | None = None,
    usar_cache: bool = True,
) -> pd.DataFrame:
    """Série mensal SGS → colunas mes (Timestamp) e valor."""
    if cache is not None and usar_cache and cache.exists() and cache.stat().st_size > 20:
        df = pd.read_csv(cache, parse_dates=["mes"])
        df["valor"] = pd.to_numeric(df["valor"], errors="coerce")
        df = df.dropna(subset=["mes"]).sort_values("mes").reset_index(drop=True)
        minimo = CACHE_MINIMO.get(codigo)
        if not df.empty and (minimo is None or df["mes"].min() <= pd.Timestamp(minimo)):
            return df

    url = BCB_SGS.format(cod=codigo)
    partes: list[pd.DataFrame] = []
    cursor = pd.Timestamp(inicio)
    fim_ts = pd.Timestamp(fim)
    while cursor <= fim_ts:
        bloco_fim = min(cursor + pd.DateOffset(years=8, months=11), fim_ts)
        dados = _http_get(
            url,
            {
                "formato": "json",
                "dataInicial": cursor.strftime("%d/%m/%Y"),
                "dataFinal": bloco_fim.strftime("%d/%m/%Y"),
            },
        )
        if dados:
            df = pd.DataFrame(dados)
            df["mes"] = pd.to_datetime(df["data"], dayfirst=True, errors="coerce")
            df["valor"] = pd.to_numeric(df["valor"], errors="coerce")
            partes.append(df[["mes", "valor"]].dropna())
        cursor = bloco_fim + pd.DateOffset(days=1)
        time.sleep(0.25)

    if not partes:
        out = pd.DataFrame(columns=["mes", "valor"])
    else:
        out = (
            pd.concat(partes, ignore_index=True)
            .drop_duplicates("mes")
            .sort_values("mes")
            .reset_index(drop=True)
        )
    if cache is not None:
        cache.parent.mkdir(parents=True, exist_ok=True)
        out.to_csv(cache, index=False)
    return out


def carregar_painel(
    pasta_cache: Path,
    *,
    usar_cache: bool = True,
    series: Iterable[Serie] = SERIES,
    inicio: date = date(1991, 1, 1),
    fim: date = date(2026, 12, 31),
    arquivos: dict[int, Path] | None = None,
) -> pd.DataFrame:
    """Painel mensal: índice = mês, colunas = código SGS.

    Se ``arquivos`` é informado, só lê esses CSV e **não** baixa do SGS
    (os testes permanecem offline).
    """
    pasta_cache.mkdir(parents=True, exist_ok=True)
    cols: dict[int, pd.Series] = {}
    for s in series:
        if arquivos is not None:
            if s.codigo not in arquivos:
                continue
            df = pd.read_csv(arquivos[s.codigo], parse_dates=["mes"])
            df["valor"] = pd.to_numeric(df["valor"], errors="coerce")
        else:
            print(f"[SGS] {s.codigo} {s.nome}")
            df = baixar_sgs(
                s.codigo,
                inicio,
                fim,
                cache=pasta_cache / f"sgs_{s.codigo}.csv",
                usar_cache=usar_cache,
            )
        if df.empty:
            continue
        ser = df.dropna(subset=["mes"]).set_index("mes")["valor"]
        ser = ser[~ser.index.duplicated(keep="last")].sort_index()
        cols[s.codigo] = ser
    if not cols:
        return pd.DataFrame()
    painel = pd.DataFrame(cols).sort_index()
    painel.index = pd.to_datetime(painel.index)
    return painel


def _nan() -> float:
    return float("nan")


def _tem(painel: pd.DataFrame, codigo: int, ts: pd.Timestamp) -> bool:
    return codigo in painel.columns and ts in painel.index and pd.notna(painel.loc[ts, codigo])


def ultimo_mes(painel: pd.DataFrame, codigo: int, ano: int) -> pd.Timestamp | None:
    if codigo not in painel.columns:
        return None
    s = painel[codigo].dropna()
    s = s[s.index.year == ano]
    if s.empty:
        return None
    return pd.Timestamp(s.index.max())


def ultimo_mes_qualquer(painel: pd.DataFrame, ano: int, codigos: Iterable[int]) -> pd.Timestamp | None:
    candidatos = [ultimo_mes(painel, c, ano) for c in codigos]
    presentes = [t for t in candidatos if t is not None]
    return max(presentes) if presentes else None


def valor_em(painel: pd.DataFrame, codigo: int, ts: pd.Timestamp | None) -> float:
    if ts is None or not _tem(painel, codigo, ts):
        return _nan()
    return float(painel.loc[ts, codigo])


def estoque_fim(painel: pd.DataFrame, codigo: int, ano: int) -> float:
    return valor_em(painel, codigo, ultimo_mes(painel, codigo, ano))


def efeito_pib(estoque_anterior_pp: float, g: float) -> float:
    """Contribuição do crescimento do PIB nominal sobre a relação dívida/PIB.

    Δ(D/Y) atribuível ao denominador = −d_{t−1} × g/(1+g),
    com g = Y_t / Y_{t−1} − 1.
    """
    if pd.isna(estoque_anterior_pp) or pd.isna(g) or g <= -1:
        return _nan()
    return float(-estoque_anterior_pp * g / (1.0 + g))


def crescimento_pib(painel: pd.DataFrame, ts_ant: pd.Timestamp, ts: pd.Timestamp) -> float:
    y0 = valor_em(painel, 4382, ts_ant)
    y1 = valor_em(painel, 4382, ts)
    if pd.isna(y0) or pd.isna(y1) or y0 == 0:
        return _nan()
    return float(y1 / y0 - 1.0)


def _ano_completo(ultimo: pd.Timestamp | None) -> bool:
    return ultimo is not None and int(ultimo.month) == 12


def nfsp_gg(
    painel: pd.DataFrame, ano: int
) -> tuple[float, float]:
    """Juros e primário do Governo Geral no ano, em p.p. do PIB.

    Ano civil fechado: séries acumuladas em 12 meses de dezembro.
    Ano incompleto (ex.: 2026*): soma dos fluxos mensais ÷ PIB de 12 meses
    do último mês — o acumulado em 12 meses **não** é o fluxo do ano.
    """
    t12 = ultimo_mes_qualquer(painel, ano, (5751, 5784, 5753, 5786))
    if (
        _ano_completo(t12)
        and _tem(painel, 5751, t12)
        and _tem(painel, 5753, t12)
        and _tem(painel, 5784, t12)
        and _tem(painel, 5786, t12)
    ):
        juros = float(painel.loc[t12, 5751] + painel.loc[t12, 5753])
        prim = float(painel.loc[t12, 5784] + painel.loc[t12, 5786])
        return juros, prim

    t_mes = ultimo_mes_qualquer(painel, ano, (4607, 4640, 4610, 4611, 4643, 4644))
    if t_mes is None:
        return _nan(), _nan()
    meses = pd.date_range(pd.Timestamp(ano, 1, 1), t_mes, freq="MS")
    # Governo Geral exige o federal no período; sem ele a soma é só regional.
    if 4607 not in painel.columns or 4640 not in painel.columns:
        return _nan(), _nan()
    fed_j = painel[4607].reindex(meses)
    fed_p = painel[4640].reindex(meses)
    if fed_j.isna().all() or fed_p.isna().all():
        return _nan(), _nan()

    def _soma(codigos: tuple[int, ...]) -> float:
        acc = 0.0
        n = 0
        for c in codigos:
            if c not in painel.columns:
                continue
            s = painel[c].reindex(meses)
            if s.notna().any():
                acc += float(s.fillna(0.0).sum())
                n += 1
        return acc if n else _nan()

    juros_rs = _soma((4607, 4610, 4611))
    prim_rs = _soma((4640, 4643, 4644))
    pib = valor_em(painel, 4382, t_mes)
    if pd.isna(pib) or pib == 0 or pd.isna(juros_rs) or pd.isna(prim_rs):
        return _nan(), _nan()
    return 100.0 * juros_rs / pib, 100.0 * prim_rs / pib


def escolher_estoque(painel: pd.DataFrame, ano: int) -> tuple[int | None, str]:
    """Série de estoque do ano.

    A metodologia atual só entra quando dá para calcular a variação
    (há dez/t−1 e o último mês de t em 13762) — a partir de 2007.
    2006 usa a metodologia antiga (4537), embora dez/2006 já seja o
    primeiro ponto da série nova.
    """
    t = ultimo_mes(painel, 13762, ano)
    t1 = ultimo_mes(painel, 13762, ano - 1)
    if t is not None and t1 is not None:
        return 13762, METODO_ATUAL
    t_old = ultimo_mes(painel, 4537, ano)
    if t_old is not None:
        return 4537, METODO_ANTIGO
    if t is not None:
        return 13762, METODO_ATUAL
    return None, ""


def _estoque_rs(painel: pd.DataFrame, ano: int, codigo_pp: int | None, ts: pd.Timestamp | None) -> float:
    if codigo_pp == 13762:
        return valor_em(painel, 13761, ts if ts is not None else ultimo_mes(painel, 13761, ano))
    if codigo_pp == 4537 and ts is not None:
        pp = valor_em(painel, 4537, ts)
        pib = valor_em(painel, 4382, ts)
        if pd.notna(pp) and pd.notna(pib):
            return pp / 100.0 * pib
    return _nan()


def linha_anual(painel: pd.DataFrame, ano: int) -> dict[str, object]:
    codigo, metodo = escolher_estoque(painel, ano)
    ts = ultimo_mes(painel, codigo, ano) if codigo else None
    ts_ant = ultimo_mes(painel, codigo, ano - 1) if codigo else None
    estoque = valor_em(painel, codigo, ts) if codigo else _nan()
    estoque_ant = valor_em(painel, codigo, ts_ant) if codigo else _nan()
    delta = (
        float(estoque - estoque_ant)
        if pd.notna(estoque) and pd.notna(estoque_ant)
        else _nan()
    )
    g = crescimento_pib(painel, ts_ant, ts) if ts is not None and ts_ant is not None else _nan()
    e_pib = efeito_pib(estoque_ant, g) if pd.notna(estoque_ant) else _nan()
    juros, prim = nfsp_gg(painel, ano)
    if pd.notna(delta) and pd.notna(juros) and pd.notna(prim) and pd.notna(e_pib):
        demais = float(delta - juros - prim - e_pib)
        variacao = float(juros + prim + e_pib + demais)
    else:
        demais = _nan()
        variacao = _nan()
    return {
        "Ano": ano,
        COL_JUROS: juros,
        COL_PRIMARIO: prim,
        COL_PIB: e_pib,
        COL_DEMAIS: demais,
        COL_VARIACAO: variacao,
        COL_ESTOQUE_PP: estoque,
        COL_ESTOQUE_RS: _estoque_rs(painel, ano, codigo, ts),
        COL_METODO: metodo,
        "ultimo_mes": ts,
        "incompleto": bool(ts is not None and ts.month < 12) or (
            ts is None and ano == ANO_FIM
        ),
    }


def anos_relatorio(painel: pd.DataFrame) -> list[int]:
    """1995–2026, cortando no último ano com qualquer observação útil."""
    if painel.empty:
        return list(range(ANO_INICIO, ANO_FIM + 1))
    ultimo = int(painel.dropna(how="all").index.max().year)
    fim = min(ANO_FIM, max(ultimo, ANO_INICIO))
    return list(range(ANO_INICIO, fim + 1))


def tabela_discriminativo(
    painel: pd.DataFrame, anos: Iterable[int] | None = None
) -> pd.DataFrame:
    if anos is None:
        anos = anos_relatorio(painel)
    return pd.DataFrame([linha_anual(painel, int(a)) for a in anos])


def tabela_anual(disc: pd.DataFrame) -> pd.DataFrame:
    """Fatores nas linhas, anos nas colunas (espelho do Discriminativo)."""
    itens = [
        (COL_JUROS, "fator"),
        (COL_PRIMARIO, "fator"),
        (COL_PIB, "fator"),
        (COL_DEMAIS, "fator"),
        (COL_VARIACAO, "total"),
        (COL_ESTOQUE_PP, "estoque"),
        (COL_ESTOQUE_RS, "estoque_rs"),
        (COL_METODO, "meta"),
    ]
    anos = [int(a) for a in disc["Ano"].tolist()]
    rows = []
    for item, papel in itens:
        rows.append(
            {
                "Item": item,
                "Papel": papel,
                **{a: disc.loc[disc["Ano"] == a, item].iloc[0] for a in anos},
            }
        )
    return pd.DataFrame(rows)


def identidade_anual(disc: pd.DataFrame) -> pd.DataFrame:
    out = disc[["Ano", COL_JUROS, COL_PRIMARIO, COL_PIB, COL_DEMAIS, COL_VARIACAO, COL_ESTOQUE_PP]].copy()
    soma = out[COL_JUROS] + out[COL_PRIMARIO] + out[COL_PIB] + out[COL_DEMAIS]
    out["soma_fatores"] = soma
    out["residuo"] = out[COL_VARIACAO] - soma
    return out


# --- Excel -----------------------------------------------------------------

AZUL = "1F4E79"
AZUL_CLARO = "D6E3F0"
DOURADO = "FFF2CC"
VERDE = "C6EFCE"
VERMELHO_TEXTO = "9B1B1B"
VERMELHO_FUNDO = "FFC7CE"
CINZA = "F2F2F2"
BRANCO = "FFFFFF"
MENOS = "\u2212"
FMT_PP = f'#,##0.00;"{MENOS}"#,##0.00;"—"'
FMT_NUM = f'#,##0.0;"{MENOS}"#,##0.0;"—"'
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def _cab(ws: Worksheet, headers: list[str], linha: int) -> None:
    fill = PatternFill("solid", fgColor=AZUL)
    font = Font(color=BRANCO, bold=True, name="Calibri", size=10)
    for col, nome in enumerate(headers, start=1):
        cell = ws.cell(linha, col, nome)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN


def _pintar_pp(cell, valor: float, *, destaque: bool = False) -> None:
    cell.number_format = FMT_PP
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        cell.value = None
        return
    cell.value = float(valor)
    if float(valor) < -1e-9:
        cell.font = Font(name="Calibri", size=9, bold=True, color=VERMELHO_TEXTO)
        cell.fill = PatternFill("solid", fgColor=VERMELHO_FUNDO)
    elif destaque:
        cell.font = Font(name="Calibri", size=9, bold=True)
    else:
        cell.font = Font(name="Calibri", size=9)


def _aba_metodologia(
    wb: Workbook,
    gerado_em: datetime,
    ultimo: pd.Timestamp,
    n_anos: int,
) -> None:
    ws = wb.active
    ws.title = "Metodologia"
    ws["A1"] = "Fatores condicionantes da Dívida Bruta do Governo Geral (1995–2026)"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=AZUL)
    ws.merge_cells("A1:B1")

    blocos = [
        ("Gerado em", gerado_em.strftime("%d/%m/%Y %H:%M")),
        ("Marker", MARKER),
        ("Fonte", "Banco Central do Brasil — SGS (Departamento de Estatísticas)"),
        ("Unidade", "Fatores e variação em pontos percentuais do PIB; estoque também em R$ milhões"),
        ("Cobertura", f"{ANO_INICIO}–{ANO_FIM}; último mês disponível: {ultimo.strftime('%m/%Y')}"),
        ("Anos na amostra", str(n_anos)),
        (
            "O que é um fator condicionante da DBGG",
            "Na Nota de Estatísticas Fiscais o Bacen decompõe a variação da "
            "relação DBGG/PIB em juros nominais apropriados, emissões líquidas "
            "de dívida, efeito da variação cambial, reconhecimento de dívidas "
            "e/ou privatizações, demais ajustes e efeito do crescimento do PIB "
            "nominal. Essa decomposição oficial não é publicada como série SGS.",
        ),
        (
            "Identidade usada aqui",
            "Δ(DBGG/PIB) = juros_NFSP_GG + primário_NFSP_GG + efeito_PIB + demais. "
            "efeito_PIB = −d_{t−1} × g/(1+g), g = PIB_12m_t / PIB_12m_{t−1} − 1 "
            "(SGS 4382). «Demais» é o resíduo que fecha a identidade.",
        ),
        (
            "Juros e primário",
            "Governo Geral = Governo Federal + estados + municípios "
            "(SGS 5751+5753 e 5784+5786 no acumulado em 12 meses; "
            "4607+4610+4611 e 4640+4643+4644 no fluxo mensal). "
            "Sinal da NFSP: necessidade positiva = déficit (aumenta a dívida); "
            "superávit primário entra negativo. A NFSP mede a dívida líquida: "
            "os juros da DBGG (bruta) são em geral maiores, e as «emissões "
            "líquidas» da Nota do Bacen não coincidem com o primário quando "
            "o Tesouro acumula ou desfaz ativos (Conta Única).",
        ),
        (
            "Quebra metodológica da DBGG",
            "Metodologia atual (SGS 13762 / 13761): estoque desde dez/2006; "
            "primeira variação anual em 2007. Metodologia até 2007 (SGS 4537): "
            "estoque desde 2002. Não se soma nem se encadeia as duas séries. "
            "Antes de 2002 o Bacen não publica DBGG oficial no SGS — as linhas "
            "1995–2001 mostram só a NFSP do Governo Geral a partir de 1999.",
        ),
        (
            "2026*",
            "Ano incompleto (último mês da amostra). Juros e primário são a "
            "soma dos fluxos mensais de janeiro até o último mês, dividida "
            "pelo PIB acumulado em 12 meses desse mês — não se usa o "
            "acumulado em 12 meses, que mistura o ano anterior.",
        ),
        (
            "Sinal e formatação",
            "Positivo = aumenta a relação DBGG/PIB. Negativos: sinal − "
            "(U+2212) em vermelho negrito sobre fundo #FFC7CE.",
        ),
        (
            "Abas",
            "Discriminativo — um ano por linha; «Variação da DBGG/PIB» = "
            "SOMA algébrica dos quatro fatores. Anual — a mesma matriz "
            "transposta. Identidade — resíduo da identidade. Grafico — "
            "contribuições anuais empilhadas.",
        ),
        (
            "Não é DLSP",
            "A Dívida Líquida do Setor Público (SGS 4513) não substitui a "
            "DBGG: o efeito cambial da DLSP inclui as reservas internacionais "
            "(sinal oposto ao da dívida externa bruta).",
        ),
    ]
    for i, (k, v) in enumerate(blocos, start=3):
        ws.cell(i, 1, k).font = Font(name="Calibri", size=10, bold=True, color=AZUL)
        ws.cell(i, 2, v).font = Font(name="Calibri", size=10)
        ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 48 if len(v) > 80 else 22
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 110


def _rotulo_ano(ano: int, incompleto: bool) -> str:
    return f"{ano}*" if incompleto else str(ano)


def _aba_discriminativo(wb: Workbook, disc: pd.DataFrame, ultimo: pd.Timestamp) -> None:
    ws = wb.create_sheet("Discriminativo", 1)
    headers = [
        "Ano",
        COL_JUROS,
        COL_PRIMARIO,
        COL_PIB,
        COL_DEMAIS,
        COL_VARIACAO,
        COL_ESTOQUE_PP,
        COL_ESTOQUE_RS,
        COL_METODO,
    ]
    ultima_col = len(headers)
    ws["A1"] = (
        "Discriminativo — fatores condicionantes da DBGG/PIB (1995–2026)"
    )
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=AZUL)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima_col)
    ws["A2"] = (
        "Cada linha é um ano. Fatores em p.p. do PIB. "
        f"A coluna «{COL_VARIACAO}» é a soma algébrica dos quatro fatores "
        "(fórmula SOMA). Positivo aumenta a dívida. "
        f"Último mês da amostra: {ultimo.strftime('%m/%Y')}. "
        "Negativos: sinal − vermelho negrito, fundo #FFC7CE."
    )
    ws["A2"].font = Font(name="Calibri", size=10, italic=True)
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima_col)
    ws.row_dimensions[2].height = 40
    _cab(ws, headers, 4)

    fill_var = PatternFill("solid", fgColor=DOURADO)
    fill_est = PatternFill("solid", fgColor=AZUL_CLARO)
    letra_ini = get_column_letter(2)
    letra_fim = get_column_letter(1 + len(FATORES))
    col_var = 2 + len(FATORES)

    for i, row in enumerate(disc.itertuples(index=False), start=5):
        valores = {disc.columns[j]: row[j] for j in range(len(disc.columns))}
        ano = int(valores["Ano"])
        incompleto = bool(valores.get("incompleto", False))
        c_ano = ws.cell(i, 1, _rotulo_ano(ano, incompleto))
        c_ano.border = THIN
        c_ano.font = Font(name="Calibri", size=10, bold=True)
        c_ano.alignment = Alignment(horizontal="center")
        if i % 2 == 0:
            c_ano.fill = PatternFill("solid", fgColor=CINZA)

        for j, nome in enumerate(FATORES, start=2):
            cell = ws.cell(i, j)
            cell.border = THIN
            _pintar_pp(cell, float(valores[nome]) if pd.notna(valores[nome]) else float("nan"))

        c_var = ws.cell(i, col_var)
        c_var.border = THIN
        _pintar_pp(
            c_var,
            float(valores[COL_VARIACAO]) if pd.notna(valores[COL_VARIACAO]) else float("nan"),
            destaque=True,
        )
        tem_fatores = all(pd.notna(valores[n]) for n in FATORES)
        if tem_fatores:
            c_var.value = f"=SUM({letra_ini}{i}:{letra_fim}{i})"
        if tem_fatores and pd.notna(valores[COL_VARIACAO]) and float(valores[COL_VARIACAO]) >= -1e-9:
            c_var.fill = fill_var
            c_var.font = Font(name="Calibri", size=9, bold=True)

        c_pp = ws.cell(i, col_var + 1)
        c_pp.border = THIN
        c_pp.fill = fill_est
        c_pp.number_format = FMT_PP
        c_pp.font = Font(name="Calibri", size=9, bold=True)
        if pd.notna(valores[COL_ESTOQUE_PP]):
            c_pp.value = float(valores[COL_ESTOQUE_PP])

        c_rs = ws.cell(i, col_var + 2)
        c_rs.border = THIN
        c_rs.fill = fill_est
        c_rs.number_format = FMT_NUM
        c_rs.font = Font(name="Calibri", size=9, bold=True)
        if pd.notna(valores[COL_ESTOQUE_RS]):
            c_rs.value = float(valores[COL_ESTOQUE_RS])

        c_met = ws.cell(i, col_var + 3, valores[COL_METODO] or None)
        c_met.border = THIN
        c_met.font = Font(name="Calibri", size=8, italic=True)
        c_met.alignment = Alignment(horizontal="center")

    tot = 5 + len(disc)
    ws.cell(tot, 1, "Soma dos fluxos").font = Font(name="Calibri", size=10, bold=True)
    ws.cell(tot, 1).fill = fill_var
    ws.cell(tot, 1).border = THIN
    primeira, ultima = 5, 4 + len(disc)
    for j in range(2, col_var + 1):
        letra = get_column_letter(j)
        cell = ws.cell(tot, j, f"=SUM({letra}{primeira}:{letra}{ultima})")
        cell.number_format = FMT_PP
        cell.font = Font(name="Calibri", size=9, bold=True)
        cell.fill = fill_var
        cell.border = THIN
    for j in range(col_var + 1, ultima_col + 1):
        ws.cell(tot, j).border = THIN
        ws.cell(tot, j).fill = fill_var

    ws.column_dimensions["A"].width = 10
    for j, nome in enumerate(headers, start=1):
        if j == 1:
            continue
        ws.column_dimensions[get_column_letter(j)].width = 16 if j < ultima_col else 18
    ws.column_dimensions[get_column_letter(ultima_col)].width = 18
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:{get_column_letter(ultima_col)}{ultima}"
    ws.row_dimensions[4].height = 36
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.cell(
        tot + 2,
        1,
        "Valores em p.p. do PIB, exceto o estoque em R$ milhões. "
        f"«{COL_VARIACAO}» = SOMA de juros + primário + efeito PIB + demais. "
        "1995–2001: sem estoque oficial de DBGG no SGS. "
        "2002–2006: metodologia até 2007 (4537). "
        "2007–2026: metodologia a partir de 2008 (13762/13761). "
        f"Coluna {ultimo.year}* = até {ultimo.strftime('%m/%Y')}.",
    ).font = Font(name="Calibri", size=8, italic=True, color="666666")


def _aba_anual(wb: Workbook, anual: pd.DataFrame, anos: list[int], ultimo: pd.Timestamp) -> None:
    ws = wb.create_sheet("Anual")
    ws["A1"] = "Fatores condicionantes da DBGG — um fator por linha"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=AZUL)
    ws["A2"] = (
        "Mesma informação do Discriminativo, transposta. "
        "p.p. do PIB, exceto estoque em R$ milhões. "
        f"Último mês: {ultimo.strftime('%m/%Y')}."
    )
    ws["A2"].font = Font(name="Calibri", size=10, italic=True)
    headers = ["Item", "Papel"] + [str(a) for a in anos]
    _cab(ws, headers, 4)
    for i, (_, row) in enumerate(anual.iterrows(), start=5):
        ws.cell(i, 1, row["Item"]).border = THIN
        ws.cell(i, 1).font = Font(name="Calibri", size=9)
        ws.cell(i, 2, row["Papel"]).border = THIN
        for j, a in enumerate(anos, start=3):
            cell = ws.cell(i, j)
            cell.border = THIN
            val = row[a]
            if row["Papel"] == "meta":
                cell.value = val if isinstance(val, str) and val else None
                cell.font = Font(name="Calibri", size=8, italic=True)
            elif row["Papel"] == "estoque_rs":
                cell.number_format = FMT_NUM
                if pd.notna(val):
                    cell.value = float(val)
                    cell.font = Font(name="Calibri", size=8)
            else:
                _pintar_pp(cell, float(val) if pd.notna(val) else float("nan"))
    ws.freeze_panes = "C5"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12
    for j in range(3, 3 + len(anos)):
        ws.column_dimensions[get_column_letter(j)].width = 9


def _aba_identidade(wb: Workbook, ident: pd.DataFrame) -> None:
    ws = wb.create_sheet("Identidade")
    ws["A1"] = "Identidade anual — Δ(DBGG/PIB) = juros + primário + efeito PIB + demais"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=AZUL)
    n = ident.dropna(subset=["residuo"])
    n_ok = int((n["residuo"].abs() <= TOLERANCIA_PP).sum()) if not n.empty else 0
    ws["A2"] = (
        f"{n_ok} de {len(n)} anos com identidade fechada "
        f"(|resíduo da soma| ≤ {TOLERANCIA_PP:.2f} p.p.). "
        "O «Demais» já é o resíduo econômico; esta coluna só confere a álgebra."
    )
    ws["A2"].font = Font(name="Calibri", size=10, italic=True)
    headers = ["Ano", COL_JUROS, COL_PRIMARIO, COL_PIB, COL_DEMAIS, "Soma", COL_VARIACAO, "Resíduo álgebra"]
    _cab(ws, headers, 4)
    for i, r in enumerate(ident.itertuples(index=False), start=5):
        vals = [int(r.Ano), r[1], r[2], r[3], r[4], r.soma_fatores, r[5], r.residuo]
        ok = pd.isna(r.residuo) or abs(float(r.residuo)) <= TOLERANCIA_PP
        for col, valor in enumerate(vals, start=1):
            cell = ws.cell(i, col, None if (isinstance(valor, float) and pd.isna(valor)) else valor)
            cell.border = THIN
            if col > 1 and not isinstance(valor, str):
                cell.number_format = FMT_PP
            if col > 1 and not ok:
                cell.fill = PatternFill("solid", fgColor="FCE4D6")
            elif ok and col > 1:
                cell.fill = PatternFill("solid", fgColor=VERDE)
            cell.font = Font(name="Calibri", size=8)
    ws.freeze_panes = "A5"
    for i, w in enumerate([10, 14, 12, 16, 14, 12, 16, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _aba_grafico(wb: Workbook, disc: pd.DataFrame) -> None:
    ws = wb.create_sheet("Grafico")
    ws["A1"] = "Contribuição anual dos fatores (p.p. do PIB)"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=AZUL)
    usable = disc.dropna(subset=[COL_VARIACAO]).copy()
    if usable.empty:
        return
    headers = ["Ano"] + list(FATORES)
    _cab(ws, headers, 3)
    for i, row in enumerate(usable.itertuples(index=False), start=4):
        valores = {usable.columns[j]: row[j] for j in range(len(usable.columns))}
        ws.cell(i, 1, int(valores["Ano"])).border = THIN
        for j, nome in enumerate(FATORES, start=2):
            cell = ws.cell(i, j, float(valores[nome]))
            cell.number_format = FMT_PP
            cell.border = THIN
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.title = "Fatores condicionantes da DBGG/PIB"
    chart.y_axis.title = "p.p. do PIB"
    chart.style = 10
    chart.height = 12
    chart.width = 22
    data = Reference(ws, min_col=2, max_col=1 + len(FATORES), min_row=3, max_row=3 + len(usable))
    cats = Reference(ws, min_col=1, min_row=4, max_row=3 + len(usable))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "A20")
    ws.column_dimensions["A"].width = 10
    for j in range(2, 6):
        ws.column_dimensions[get_column_letter(j)].width = 18


def escrever_planilha(
    *,
    disc: pd.DataFrame,
    anual: pd.DataFrame,
    ident: pd.DataFrame,
    saida: Path,
    ultimo: pd.Timestamp,
) -> Path:
    saida.parent.mkdir(parents=True, exist_ok=True)
    anos = [int(a) for a in disc["Ano"].tolist()]
    wb = Workbook()
    _aba_metodologia(wb, datetime.now(), ultimo, len(disc))
    _aba_discriminativo(wb, disc, ultimo)
    _aba_anual(wb, anual, anos, ultimo)
    _aba_identidade(wb, ident)
    _aba_grafico(wb, disc)
    wb.save(saida)
    print(f"[OK] Planilha: {saida} ({saida.stat().st_size / 1024:.1f} KB)")
    return saida


def processar(
    *,
    pasta_cache: Path,
    saida: Path,
    usar_cache: bool = True,
    arquivos: dict[int, Path] | None = None,
) -> Path:
    print(f"[{MARKER}]")
    painel = carregar_painel(pasta_cache, usar_cache=usar_cache, arquivos=arquivos)
    if painel.empty:
        raise RuntimeError("painel vazio — sem séries SGS")
    anos = anos_relatorio(painel)
    ultimo = painel.dropna(how="all").index.max()
    print(f"[INFO] painel {painel.index.min().date()} → {ultimo.date()} | anos={anos[0]}–{anos[-1]}")
    disc = tabela_discriminativo(painel, anos)
    anual = tabela_anual(disc)
    ident = identidade_anual(disc)
    path = escrever_planilha(
        disc=disc, anual=anual, ident=ident, saida=saida, ultimo=ultimo
    )
    print("\n=== Discriminativo anual (p.p. do PIB) ===")
    for _, r in disc.iterrows():
        ano = int(r["Ano"])
        if ano not in (1999, 2002, 2006, 2007, 2015, 2020, 2024, 2025, 2026):
            continue
        rot = _rotulo_ano(ano, bool(r["incompleto"]))
        def _f(x: object) -> str:
            return "     n/d" if x is None or (isinstance(x, float) and pd.isna(x)) else f"{float(x):8.2f}"

        print(
            f"  {rot:>6}  juros={_f(r[COL_JUROS])}  prim={_f(r[COL_PRIMARIO])}  "
            f"PIB={_f(r[COL_PIB])}  demais={_f(r[COL_DEMAIS])}  "
            f"Δ={_f(r[COL_VARIACAO])}  est={_f(r[COL_ESTOQUE_PP])}  {r[COL_METODO]}"
        )
    return path


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--pasta-cache", type=Path, default=ROOT / "data" / "sgs")
    p.add_argument(
        "--saida",
        type=Path,
        default=ROOT / "output" / "fatores_condicionantes_dbgg_1995_2026.xlsx",
    )
    p.add_argument("--sem-cache", action="store_true")
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        processar(
            pasta_cache=args.pasta_cache,
            saida=args.saida,
            usar_cache=not args.sem_cache,
        )
    except Exception as exc:  # noqa: BLE001
        print(f"ERRO: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
