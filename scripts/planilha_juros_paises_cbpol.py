#!/usr/bin/env python3
"""Taxas básicas de juros diárias por país (BIS WS_CBPOL).

Lê o export largo ``WS_CBPOL_csv_col.csv`` e grava um Excel com uma aba por
área de referência. Cada aba traz a taxa diária de 01/01/1995 a 31/08/2026,
sem as datas que caem em sábado ou domingo.

Uso:
  python3 scripts/planilha_juros_paises_cbpol.py --arquivo WS_CBPOL_csv_col.csv
  python3 scripts/planilha_juros_paises_cbpol.py --baixar
"""

from __future__ import annotations

import argparse
import csv
import sys
import urllib.request
import zipfile
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
DATA_INICIO = date(1995, 1, 1)
DATA_FIM = date(2026, 8, 31)
BIS_ZIP_URL = "https://data.bis.org/static/bulk/WS_CBPOL_csv_col.zip"
CAMINHO_WIN = Path(
    r"C:\Arquivos de Programas RFB\ContAgilAppBeta64\python_jep\winpython"
    r"\bis-mcp\WS_CBPOL_csv_col.csv"
)
SAIDA_PADRAO = ROOT / "output" / "juros_basicos_diarios_paises_1995_2026.xlsx"
ABA_INDICE = "Indice"
FMT_DATA = "DD/MM/YYYY"
FMT_TAXA = "0.00##"
CARACTERES_ABA = set(r"[]:*?/\\")


@dataclass(frozen=True)
class SeriePais:
    codigo: str
    pais: str
    definicao: str
    observacoes: tuple[tuple[date, float], ...]
    fins_de_semana_omitidos: int

    @property
    def aba(self) -> str:
        return self.pais


def parse_data_coluna(nome: str) -> date | None:
    """Data de coluna diária ``YYYY-MM-DD``. Colunas mensais ``YYYY-MM`` ficam de fora."""
    if len(nome) != 10 or nome[4] != "-" or nome[7] != "-":
        return None
    try:
        return date(int(nome[0:4]), int(nome[5:7]), int(nome[8:10]))
    except ValueError:
        return None


def valor_taxa(texto: str | None) -> float | None:
    if texto is None:
        return None
    s = texto.strip()
    if not s or s.lower() == "nan":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def classificar_colunas(
    header: list[str], inicio: date, fim: date
) -> tuple[list[tuple[int, date]], list[int]]:
    """Índices das colunas diárias no intervalo: dias de semana e fins de semana."""
    uteis: list[tuple[int, date]] = []
    fins: list[int] = []
    for i, nome in enumerate(header):
        dt = parse_data_coluna(nome)
        if dt is None or dt < inicio or dt > fim:
            continue
        if dt.weekday() >= 5:
            fins.append(i)
        else:
            uteis.append((i, dt))
    return uteis, fins


def nome_aba(pais: str, usados: set[str]) -> str:
    limpo = "".join(c for c in pais if c not in CARACTERES_ABA).strip() or "Pais"
    base = limpo[:31]
    nome = base
    n = 2
    while nome.casefold() in usados:
        sufixo = f" {n}"
        nome = base[: 31 - len(sufixo)] + sufixo
        n += 1
    usados.add(nome.casefold())
    return nome


def extrair_series(
    caminho: Path,
    inicio: date = DATA_INICIO,
    fim: date = DATA_FIM,
) -> list[SeriePais]:
    """Séries diárias (FREQ=D), só dias de semana com taxa numérica no intervalo."""
    series: list[SeriePais] = []
    with caminho.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.reader(handle)
        try:
            header = next(reader)
        except StopIteration as exc:
            raise ValueError(f"CSV vazio: {caminho}") from exc
        uteis, fins = classificar_colunas(header, inicio, fim)
        if not uteis:
            raise ValueError(
                "Nenhuma coluna diária no período "
                f"{inicio.isoformat()} a {fim.isoformat()}."
            )
        for row in reader:
            if not row or row[0].strip() != "D":
                continue
            obs: list[tuple[date, float]] = []
            for i, dt in uteis:
                if i >= len(row):
                    break
                taxa = valor_taxa(row[i])
                if taxa is not None:
                    obs.append((dt, taxa))
            omitidos = 0
            for i in fins:
                if i >= len(row):
                    break
                if valor_taxa(row[i]) is not None:
                    omitidos += 1
            series.append(
                SeriePais(
                    codigo=row[2].strip() if len(row) > 2 else "",
                    pais=row[3].strip() if len(row) > 3 else "",
                    definicao=row[6].strip() if len(row) > 6 else "",
                    observacoes=tuple(obs),
                    fins_de_semana_omitidos=omitidos,
                )
            )
    series.sort(key=lambda s: (s.pais.casefold(), s.codigo))
    return series


def _celula(ws, valor, formato: str | None = None, negrito: bool = False):
    cel = WriteOnlyCell(ws, value=valor)
    if formato:
        cel.number_format = formato
    if negrito:
        cel.font = Font(bold=True, name="Calibri")
    return cel


def gravar_planilha(series: list[SeriePais], destino: Path) -> dict[str, str]:
    """Grava o índice e uma aba por país. Devolve código → nome da aba."""
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=True)
    wb.properties.creator = "planilha_juros_paises_cbpol"
    wb.properties.title = "Taxas básicas de juros diárias (BIS WS_CBPOL)"

    indice = wb.create_sheet(ABA_INDICE)
    indice.freeze_panes = "A3"
    for letra, largura in (
        ("A", 12),
        ("B", 22),
        ("C", 22),
        ("D", 16),
        ("E", 16),
        ("F", 16),
        ("G", 22),
        ("H", 22),
        ("I", 28),
        ("J", 80),
    ):
        indice.column_dimensions[letra].width = largura

    titulo = (
        "Taxas básicas de juros diárias (BIS WS_CBPOL), "
        f"{DATA_INICIO.strftime('%d/%m/%Y')} a {DATA_FIM.strftime('%d/%m/%Y')}. "
        "Sábados e domingos excluídos. Unidade: % ao ano, fim de período."
    )
    indice.append([_celula(indice, titulo, negrito=True)])
    cab_ind = [
        "Código",
        "País",
        "Aba",
        "Observações",
        "Primeira data",
        "Última data",
        "Taxa inicial (% a.a.)",
        "Taxa final (% a.a.)",
        "Sábados e domingos omitidos",
        "Definição BIS",
    ]
    indice.append([_celula(indice, h, negrito=True) for h in cab_ind])

    usados: set[str] = {ABA_INDICE.casefold()}
    abas: dict[str, str] = {}
    blocos: list[tuple[SeriePais, str]] = []
    for serie in series:
        aba = nome_aba(serie.pais or serie.codigo or "Pais", usados)
        abas[serie.codigo] = aba
        blocos.append((serie, aba))
        if serie.observacoes:
            primeira, taxa0 = serie.observacoes[0]
            ultima, taxan = serie.observacoes[-1]
        else:
            primeira = ultima = None
            taxa0 = taxan = None
        indice.append(
            [
                serie.codigo,
                serie.pais,
                aba,
                len(serie.observacoes),
                _celula(indice, primeira, FMT_DATA) if primeira else None,
                _celula(indice, ultima, FMT_DATA) if ultima else None,
                _celula(indice, taxa0, FMT_TAXA) if taxa0 is not None else None,
                _celula(indice, taxan, FMT_TAXA) if taxan is not None else None,
                serie.fins_de_semana_omitidos,
                serie.definicao,
            ]
        )

    for serie, aba in blocos:
        ws = wb.create_sheet(aba)
        ws.freeze_panes = "A2"
        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 22
        ws.append(
            [
                _celula(ws, "Data", negrito=True),
                _celula(ws, "Taxa básica (% a.a.)", negrito=True),
            ]
        )
        for dt, taxa in serie.observacoes:
            ws.append(
                [
                    _celula(ws, dt, FMT_DATA),
                    _celula(ws, taxa, FMT_TAXA),
                ]
            )
        if serie.observacoes:
            ws.auto_filter.ref = f"A1:B{len(serie.observacoes) + 1}"

    wb.save(destino)
    return abas


def baixar_csv(destino: Path) -> Path:
    destino.parent.mkdir(parents=True, exist_ok=True)
    zip_path = destino.with_suffix(".zip")
    urllib.request.urlretrieve(BIS_ZIP_URL, zip_path)
    with zipfile.ZipFile(zip_path) as arquivo:
        nome = next(n for n in arquivo.namelist() if n.lower().endswith(".csv"))
        arquivo.extract(nome, destino.parent)
        extraido = destino.parent / nome
    if extraido.resolve() != destino.resolve():
        extraido.replace(destino)
    return destino


def resolver_arquivo(explicito: Path | None, baixar: bool) -> Path:
    if explicito is not None:
        if not explicito.is_file():
            raise FileNotFoundError(explicito)
        return explicito
    candidatos = [
        Path.cwd() / "WS_CBPOL_csv_col.csv",
        CAMINHO_WIN,
        ROOT / "data" / "WS_CBPOL_csv_col.csv",
        ROOT / "data" / "raw" / "WS_CBPOL_csv_col.csv",
        ROOT / "attachments" / "WS_CBPOL_csv_col.csv",
    ]
    for caminho in candidatos:
        if caminho.is_file():
            return caminho
    if baixar:
        return baixar_csv(ROOT / "data" / "raw" / "WS_CBPOL_csv_col.csv")
    raise FileNotFoundError(
        "WS_CBPOL_csv_col.csv não encontrado. Passe --arquivo ou use --baixar."
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--arquivo", type=Path, default=None, help="WS_CBPOL_csv_col.csv")
    parser.add_argument("--saida", type=Path, default=SAIDA_PADRAO)
    parser.add_argument(
        "--baixar",
        action="store_true",
        help="Baixa o ZIP oficial do BIS se o CSV local não existir.",
    )
    args = parser.parse_args(argv)
    origem = resolver_arquivo(args.arquivo, args.baixar)
    print(f"Lendo {origem}")
    series = extrair_series(origem)
    abas = gravar_planilha(series, args.saida)
    obs = sum(len(s.observacoes) for s in series)
    omitidos = sum(s.fins_de_semana_omitidos for s in series)
    print(f"Países/áreas: {len(series)} | abas: {len(abas) + 1} (com {ABA_INDICE})")
    print(f"Dias de semana com taxa: {obs}")
    print(f"Sábados e domingos omitidos (tinham taxa): {omitidos}")
    print(f"Planilha: {args.saida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
