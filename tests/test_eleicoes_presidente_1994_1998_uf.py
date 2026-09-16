"""Testes da planilha presidencial 1994 e 1998 (1º turno por UF)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from scripts.eleicoes_presidente_1994_1998_uf import (
    CACHE,
    carregar_dados,
    chave_candidato,
    conferir_totais,
    dois_mais_votados,
    gravar_xlsx,
    linha_uf,
    montar_tabelas,
    nome_curto,
    parse_pagina,
    votos_de,
)


HTML_AC_1994 = """
<B>Presidential Election Results - Acre</B>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Registered Electors&nbsp;</TD>
  <TD>&nbsp;&nbsp;263,162</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Voters&nbsp;</TD>
  <TD>&nbsp;&nbsp;206,393</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Blank Votes&nbsp;</TD>
  <TD>&nbsp;&nbsp;23,843</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Invalid Votes&nbsp;</TD>
  <TD>&nbsp;&nbsp;15,665</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Valid Votes&nbsp;</TD>
  <TD>&nbsp;&nbsp;166,885</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Fernando Henrique Cardoso (PSDB)&nbsp;</TD>
  <TD>&nbsp;&nbsp;90,132</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Luiz In&aacute;cio Lula da Silva (PT)&nbsp;</TD>
  <TD>&nbsp;&nbsp;39,656</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;En&eacute;as Ferreira Carneiro (PRONA)&nbsp;</TD>
  <TD>&nbsp;&nbsp;12,636</TD>
</TR>
"""


def _cand(nome: str, votos: int) -> dict:
    return {"nome": nome, "votos": votos}


def _uf(fhc: int, lula: int, extra: list[dict] | None = None, validos: int | None = None) -> dict:
    cands = [
        _cand("Fernando Henrique Cardoso (PSDB)", fhc),
        _cand("Luiz Inácio Lula da Silva (PT)", lula),
    ]
    if extra:
        cands.extend(extra)
    if validos is None:
        validos = sum(c["votos"] for c in cands)
    return {"unidade": "X", "validos": validos, "candidatos": cands}


def test_parse_pagina_acre():
    parsed = parse_pagina(HTML_AC_1994)
    assert parsed["unidade"] == "Acre"
    assert parsed["validos"] == 166_885
    assert votos_de(parsed["candidatos"], "fhc") == 90_132
    assert votos_de(parsed["candidatos"], "lula") == 39_656


def test_nome_curto_e_chave():
    assert nome_curto("Fernando Henrique Cardoso (PSDB)") == "FHC"
    assert nome_curto("Luiz Inácio Lula da Silva (PT)") == "Lula"
    assert nome_curto("Ciro Ferreira Gomes (PL / PPS / PAN)") == "Ciro Gomes"
    assert chave_candidato("Fernando Henrique Cardoso (PPB / PTB / PFL / PSD / PSDB)") == "fhc"


def test_dois_mais_votados():
    br = _uf(34_364_961, 17_122_127, extra=[_cand("Enéas Ferreira Carneiro (PRONA)", 4_671_457)])
    c1, c2 = dois_mais_votados(br)
    assert c1.chave == "fhc" and c1.nome == "FHC"
    assert c2.chave == "lula" and c2.nome == "Lula"


def test_montar_tabelas_e_ceara_1998():
    from scripts import eleicoes_presidente_1994_1998_uf as mod

    bruto = {
        "1994": {
            "SP": _uf(100, 50),
            "CE": _uf(80, 40),
            "ZZ": _uf(3, 2),
            "BR": _uf(183, 92),
        },
        "1998": {
            "SP": _uf(110, 60),
            "CE": _uf(
                80,
                90,
                extra=[_cand("Ciro Ferreira Gomes (PPS)", 100)],
            ),
            "ZZ": _uf(4, 3),
            "BR": _uf(194, 153, extra=[_cand("Ciro Ferreira Gomes (PPS)", 100)]),
        },
    }
    orig = list(mod.UFS)
    try:
        mod.UFS[:] = ["SP", "CE"]
        c1, c2 = dois_mais_votados(bruto["1994"]["BR"])
        ufs, exterior, extra = montar_tabelas(bruto, c1, c2)
        assert list(ufs["uf"]) == ["SP", "CE"]
        ce = ufs.loc[ufs["uf"] == "CE"].iloc[0]
        assert ce["y1998_vencedor"] == "Ciro Gomes"
        assert int(ce["y1998_c1_votos"]) == 80
        assert int(ce["y1998_c2_votos"]) == 90
        assert list(exterior["uf"]) == ["ZZ"]
        total = extra.loc[extra["uf"] == "UFs"].iloc[0]
        assert int(total["y1994_c1_votos"]) == 180
    finally:
        mod.UFS[:] = orig


def test_gravar_xlsx(tmp_path: Path):
    from scripts import eleicoes_presidente_1994_1998_uf as mod

    bruto = {
        "1994": {
            "SP": _uf(100, 50),
            "CE": _uf(80, 40),
            "ZZ": _uf(3, 2),
            "BR": _uf(183, 92),
        },
        "1998": {
            "SP": _uf(110, 60),
            "CE": _uf(80, 70),
            "ZZ": _uf(4, 3),
            "BR": _uf(194, 133),
        },
    }
    orig = list(mod.UFS)
    try:
        mod.UFS[:] = ["SP", "CE"]
        c1, c2 = dois_mais_votados(bruto["1994"]["BR"])
        ufs, exterior, extra = montar_tabelas(bruto, c1, c2)
        path = tmp_path / "e.xlsx"
        gravar_xlsx(ufs, exterior, extra, c1, c2, path)
        wb = load_workbook(path)
        assert wb.sheetnames == ["Por_UF", "Brasil_e_exterior", "Fonte"]
        ws = wb["Por_UF"]
        assert "1994" in str(ws["D1"].value)
        assert "1998" in str(ws["J1"].value)
        assert ws["D2"].value == "FHC (votos válidos)"
        assert ws["F2"].value == "Lula (votos válidos)"
        ufs_aba = {ws.cell(row, 1).value for row in range(3, ws.max_row + 1)}
        assert {"SP", "CE", "UFs"} <= ufs_aba
        assert "ZZ" not in ufs_aba
    finally:
        mod.UFS[:] = orig


def test_cache_oficial_totais():
    if not CACHE.exists():
        return
    bruto = carregar_dados()
    conferir_totais(bruto)
    c1, c2 = dois_mais_votados(bruto["1994"]["BR"])
    assert c1.chave == "fhc"
    assert c2.chave == "lula"
    ce = linha_uf("CE", bruto, c1, c2)
    assert ce["y1998_vencedor"] == "Ciro Gomes"
    rs = linha_uf("RS", bruto, c1, c2)
    assert rs["y1994_vencedor"] == "Lula"
    assert int(rs["y1994_c2_votos"]) == 1_610_379
