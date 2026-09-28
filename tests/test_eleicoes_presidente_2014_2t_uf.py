"""Testes da planilha presidencial 2014 (2º turno por UF)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from scripts.eleicoes_presidente_2014_2t_uf import (
    CACHE,
    carregar_dados,
    chave_candidato,
    conferir_totais,
    dois_mais_votados,
    gravar_xlsx,
    linha_uf,
    montar_tabelas,
    nome_curto,
    parse_pagina,
    votos_de,
)


HTML_MG = """
<B>October 5, 2014 Presidential Election Results - Minas Gerais</B>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Valid Votes&nbsp;</TD>
  <TD>&nbsp;&nbsp;11,106,883</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Dilma (PT - PT / PMDB)&nbsp;</TD>
  <TD>&nbsp;&nbsp;4,829,513</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;A&eacute;cio Neves (PSDB - PSDB)&nbsp;</TD>
  <TD>&nbsp;&nbsp;4,414,452</TD>
</TR>
<TR>
  <TD COLSPAN=3><B>October 26, 2014 Runoff Election Results - Minas Gerais</B></TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Valid Votes&nbsp;</TD>
  <TD>&nbsp;&nbsp;11,408,243</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;Dilma (PT - PT / PMDB)&nbsp;</TD>
  <TD>&nbsp;&nbsp;5,979,422</TD>
</TR>
<TR ALIGN=RIGHT>
  <TD ALIGN=LEFT NOWRAP>&nbsp;A&eacute;cio Neves (PSDB - PSDB)&nbsp;</TD>
  <TD>&nbsp;&nbsp;5,428,821</TD>
</TR>
"""


def _cand(nome: str, votos: int) -> dict:
    return {"nome": nome, "votos": votos}


def _uf(dilma: int, aecio: int, validos: int | None = None) -> dict:
    cands = [
        _cand("Dilma (PT - PT / PMDB)", dilma),
        _cand("Aécio Neves (PSDB - PSDB)", aecio),
    ]
    return {
        "1t": {"validos": 1, "candidatos": cands},
        "2t": {
            "unidade": "X",
            "validos": dilma + aecio if validos is None else validos,
            "candidatos": cands,
        },
    }


def test_parse_pagina_separa_turnos():
    parsed = parse_pagina(HTML_MG)
    assert parsed["2t"]["validos"] == 11_408_243
    assert votos_de(parsed["2t"]["candidatos"], "dilma") == 5_979_422
    assert votos_de(parsed["2t"]["candidatos"], "aecio") == 5_428_821
    assert parsed["1t"]["validos"] == 11_106_883


def test_nome_curto_e_chave():
    assert nome_curto("Dilma (PT - PT / PMDB / PSD)") == "Dilma"
    assert nome_curto("Aécio Neves (PSDB - PSDB)") == "Aécio Neves"
    assert chave_candidato("Aécio Neves (PSDB - PSDB / DEM)") == "aecio"
    assert chave_candidato("Dilma (PT)") == "dilma"


def test_dois_mais_votados():
    br = _uf(54_501_118, 51_041_155)["2t"]
    c1, c2 = dois_mais_votados(br)
    assert c1.chave == "dilma" and c1.nome == "Dilma"
    assert c2.chave == "aecio" and c2.nome == "Aécio Neves"


def test_montar_tabelas_separa_exterior():
    from scripts import eleicoes_presidente_2014_2t_uf as mod

    bruto = {
        "SP": _uf(100, 200),
        "MG": _uf(150, 140),
        "ZZ": _uf(10, 20),
        "BR": _uf(360, 260),
    }
    orig = list(mod.UFS)
    try:
        mod.UFS[:] = ["SP", "MG"]
        c1, c2 = dois_mais_votados(bruto["BR"]["2t"])
        assert c1.chave == "dilma" and c2.chave == "aecio"
        ufs, exterior, extra = montar_tabelas(bruto, c1, c2)
        assert list(ufs["uf"]) == ["SP", "MG"]
        assert ufs.loc[ufs["uf"] == "SP", "vencedor"].iloc[0] == "Aécio Neves"
        assert ufs.loc[ufs["uf"] == "MG", "vencedor"].iloc[0] == "Dilma"
        assert list(exterior["uf"]) == ["ZZ"]
        total = extra.loc[extra["uf"] == "UFs"].iloc[0]
        assert int(total["c1_votos"]) == 250
        assert int(total["c2_votos"]) == 340
    finally:
        mod.UFS[:] = orig


def test_gravar_xlsx(tmp_path: Path):
    from scripts import eleicoes_presidente_2014_2t_uf as mod

    bruto = {
        "SP": _uf(100, 200),
        "MG": _uf(150, 140),
        "ZZ": _uf(10, 20),
        "BR": _uf(360, 260),
    }
    orig = list(mod.UFS)
    try:
        mod.UFS[:] = ["SP", "MG"]
        c1, c2 = dois_mais_votados(bruto["BR"]["2t"])
        ufs, exterior, extra = montar_tabelas(bruto, c1, c2)
        path = tmp_path / "e.xlsx"
        gravar_xlsx(ufs, exterior, extra, c1, c2, path)
        wb = load_workbook(path)
        assert wb.sheetnames == ["Por_UF", "Brasil_e_exterior", "Fonte"]
        ws = wb["Por_UF"]
        assert "2º turno" in str(ws["D1"].value)
        assert ws["D2"].value == "Dilma (votos válidos)"
        assert ws["F2"].value == "Aécio Neves (votos válidos)"
        ufs_aba = {ws.cell(row, 1).value for row in range(3, ws.max_row + 1)}
        assert {"SP", "MG", "UFs"} <= ufs_aba
        assert "ZZ" not in ufs_aba
    finally:
        mod.UFS[:] = orig


def test_cache_oficial_totais():
    if not CACHE.exists():
        return
    bruto = carregar_dados()
    conferir_totais(bruto)
    c1, c2 = dois_mais_votados(bruto["BR"]["2t"])
    assert c1.chave == "dilma"
    assert c2.chave == "aecio"
    mg = linha_uf("MG", bruto, c1, c2)
    assert mg["vencedor"] == "Dilma"
    assert int(mg["c1_votos"]) == 5_979_422
    sp = linha_uf("SP", bruto, c1, c2)
    assert sp["vencedor"] == "Aécio Neves"
    assert int(sp["c2_votos"]) == 15_296_289
