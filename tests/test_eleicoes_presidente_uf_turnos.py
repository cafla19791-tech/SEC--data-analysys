"""Testes da planilha presidencial 2022 (1º e 2º turnos por UF)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from scripts.eleicoes_presidente_uf_turnos import (
    Candidato,
    conferir_soma_ufs_mais_exterior,
    conferir_totais_nacionais,
    dois_mais_votados,
    fundir_turnos,
    gravar_xlsx,
    montar_tabelas,
    parse_abrangencia,
    pct_oficial,
    url_turno,
    votos_int,
)


def _cand(n: str, nome: str, vap: int, pvap: str, seq: str = "1") -> dict:
    return {
        "n": n,
        "nm": nome,
        "cc": "PARTIDO",
        "vap": str(vap),
        "pvap": pvap,
        "seq": seq,
    }


def _payload(uf: str, c1: int, c2: int, extra: list[dict] | None = None) -> dict:
    validos = c1 + c2 + sum(int(c["vap"]) for c in (extra or []))
    cand = [
        _cand("13", "LULA", c1, "50,20"),
        _cand("22", "JAIR BOLSONARO", c2, "49,80"),
    ]
    if extra:
        cand.extend(extra)
    return {
        "cdabr": uf,
        "vv": str(validos),
        "vb": "10",
        "vn": "20",
        "c": str(validos + 30),
        "a": "5",
        "e": str(validos + 35),
        "pst": "100,00",
        "cand": cand,
    }


C1 = Candidato("13", "LULA", "Lula", "PT", 57_259_504, 48.43)
C2 = Candidato("22", "JAIR BOLSONARO", "Jair Bolsonaro", "PL", 51_072_345, 43.20)


def test_pct_oficial_virgula_e_milhar():
    assert pct_oficial("50,90") == 50.90
    assert pct_oficial("1.620,97") == 1620.97


def test_votos_int_aceita_string_e_int():
    assert votos_int("57259504") == 57_259_504
    assert votos_int(51072345) == 51_072_345


def test_url_turno_padding():
    assert url_turno(544, "sp").endswith("/sp/sp-c0001-e000544-r.json")
    assert "/ele2022/545/" in url_turno(545, "BR")


def test_dois_mais_votados_ordena_por_votos():
    payload = _payload(
        "br",
        57_259_504,
        51_072_345,
        extra=[_cand("15", "SIMONE TEBET", 4_915_423, "4,16", "3")],
    )
    c1, c2 = dois_mais_votados(payload)
    assert c1.numero == "13" and c1.nome == "Lula"
    assert c2.numero == "22" and c2.nome == "Jair Bolsonaro"
    assert c1.votos_brasil_1t == 57_259_504


def test_parse_e_fundir_vencedor():
    t1 = parse_abrangencia(_payload("mg", 6_190_960, 6_141_310), C1, C2)
    t2 = parse_abrangencia(_payload("mg", 6_313_684, 6_165_150), C1, C2)
    row = fundir_turnos(t1, t2)
    assert row["uf"] == "MG"
    assert row["unidade"] == "Minas Gerais"
    assert row["t1_vencedor"] == "Lula"
    assert row["t2_c1_votos"] == 6_313_684
    assert row["t1_votos_validos"] == 6_190_960 + 6_141_310


def test_montar_tabelas_separa_exterior_e_soma_ufs():
    linhas = [
        fundir_turnos(
            parse_abrangencia(_payload("sp", 100, 200), C1, C2),
            parse_abrangencia(_payload("sp", 110, 220), C1, C2),
        ),
        fundir_turnos(
            parse_abrangencia(_payload("df", 40, 60), C1, C2),
            parse_abrangencia(_payload("df", 45, 70), C1, C2),
        ),
        fundir_turnos(
            parse_abrangencia(_payload("zz", 10, 8), C1, C2),
            parse_abrangencia(_payload("zz", 12, 9), C1, C2),
        ),
    ]
    brasil = fundir_turnos(
        parse_abrangencia(_payload("br", 150, 268), C1, C2),
        parse_abrangencia(_payload("br", 167, 299), C1, C2),
    )
    ufs, exterior, extra = montar_tabelas(linhas, brasil, C1, C2)
    assert list(ufs["uf"]) == ["SP", "DF"]
    assert list(exterior["uf"]) == ["ZZ"]
    total = extra.loc[extra["uf"] == "UFs"].iloc[0]
    assert int(total["t1_c1_votos"]) == 140
    assert int(total["t1_c2_votos"]) == 260
    assert int(total["t2_c1_votos"]) == 155
    assert "ZZ" not in set(ufs["uf"])
    br = extra.loc[extra["uf"] == "BR"].iloc[0]
    assert int(br["t1_votos_validos"]) == 418


def test_conferir_soma_ufs_mais_exterior():
    linhas = [
        fundir_turnos(
            parse_abrangencia(_payload("sp", 100, 200), C1, C2),
            parse_abrangencia(_payload("sp", 110, 220), C1, C2),
        ),
        fundir_turnos(
            parse_abrangencia(_payload("zz", 10, 8), C1, C2),
            parse_abrangencia(_payload("zz", 12, 9), C1, C2),
        ),
    ]
    brasil = fundir_turnos(
        parse_abrangencia(_payload("br", 110, 208), C1, C2),
        parse_abrangencia(_payload("br", 122, 229), C1, C2),
    )
    ufs, exterior, extra = montar_tabelas(linhas, brasil, C1, C2)
    conferir_soma_ufs_mais_exterior(ufs, exterior, extra.loc[extra["uf"] == "BR"].iloc[0])


def test_conferir_totais_nacionais_ok():
    br_1t = {
        "vv": str(TSE_VV_1T := 118_229_719),
        "cand": [
            _cand("13", "LULA", 57_259_504, "48,43"),
            _cand("22", "JAIR BOLSONARO", 51_072_345, "43,20"),
        ],
    }
    br_2t = {
        "vv": "118552353",
        "cand": [
            _cand("13", "LULA", 60_345_999, "50,90"),
            _cand("22", "JAIR BOLSONARO", 58_206_354, "49,10"),
        ],
    }
    conferir_totais_nacionais(C1, C2, br_1t, br_2t)
    assert TSE_VV_1T == 118_229_719


def test_gravar_xlsx(tmp_path: Path):
    linhas = [
        fundir_turnos(
            parse_abrangencia(_payload("sp", 100, 200), C1, C2),
            parse_abrangencia(_payload("sp", 110, 220), C1, C2),
        ),
        fundir_turnos(
            parse_abrangencia(_payload("df", 40, 60), C1, C2),
            parse_abrangencia(_payload("df", 45, 70), C1, C2),
        ),
        fundir_turnos(
            parse_abrangencia(_payload("zz", 10, 8), C1, C2),
            parse_abrangencia(_payload("zz", 12, 9), C1, C2),
        ),
    ]
    brasil = fundir_turnos(
        parse_abrangencia(_payload("br", 150, 268), C1, C2),
        parse_abrangencia(_payload("br", 167, 299), C1, C2),
    )
    ufs, exterior, extra = montar_tabelas(linhas, brasil, C1, C2)
    path = tmp_path / "eleicoes.xlsx"
    gravar_xlsx(ufs, exterior, extra, C1, C2, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Por_UF", "Brasil_e_exterior", "Fonte"]
    ws = wb["Por_UF"]
    assert ws["A1"].value == "Unidade da Federação"
    assert "1º turno" in str(ws["D1"].value)
    assert "2º turno" in str(ws["J1"].value)
    assert ws["D2"].value == "Lula (votos válidos)"
    assert ws["J2"].value == "Lula (votos válidos)"
    assert ws["F2"].value == "Jair Bolsonaro (votos válidos)"
    ufs_na_aba = {ws.cell(row, 1).value for row in range(3, ws.max_row + 1)}
    assert {"SP", "DF", "UFs"} <= ufs_na_aba
    assert "ZZ" not in ufs_na_aba
    assert ws.max_row == 5  # 2 UFs + total
    sp = next(ws.cell(row, 4).value for row in range(3, ws.max_row) if ws.cell(row, 1).value == "SP")
    assert sp == 100
    wb2 = wb["Brasil_e_exterior"]
    assert {wb2.cell(row, 1).value for row in range(3, wb2.max_row + 1)} == {"BR", "ZZ", "UFs"}
    assert "TSE" in str(wb["Fonte"]["A3"].value)
