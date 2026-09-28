"""Planilha de lucros líquidos atualizados pelo IPCA até 31/08/2026."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from scripts.lucros_bancos_ipca import (
    BRADESCO,
    ITAU,
    ITAU_UNIBANCO,
    SEMESTRE_2026,
    UNIBANCO,
    fator_ipca,
    gravar_planilha,
    montar_tabela,
)


def _ipca_sintetico() -> pd.DataFrame:
    meses = pd.period_range("2001-12", "2026-08", freq="M").to_timestamp("M")
    # 1% ao mês, estável, para o fator ser auditável.
    df = pd.DataFrame({"mes": meses, "ipca_pct": 1.0})
    df.loc[df["mes"].dt.to_period("M") == "2001-12", "ipca_pct"] = 0.0
    df["indice"] = (1.0 + df["ipca_pct"] / 100.0).cumprod()
    return df.reset_index(drop=True)


def test_fator_aplica_so_os_meses_depois_da_data_base():
    ipca = _ipca_sintetico()
    # De 31/01/2002 a 28/02/2002 entra apenas o IPCA de fevereiro (1%).
    assert fator_ipca(ipca, pd.Timestamp("2002-01-31"), pd.Timestamp("2002-02-28")) == pytest.approx(1.01)
    # De 31/12/2002 a 31/08/2026: janeiro/2003 até agosto/2026.
    meses = (2026 - 2003) * 12 + 8  # 2003–2025 completos + jan–ago/2026
    assert meses == 284
    assert fator_ipca(ipca, pd.Timestamp("2002-12-31")) == pytest.approx(1.01**284)


def test_semestre_de_2026_usa_junho_como_base():
    ipca = _ipca_sintetico()
    tabela = montar_tabela(ipca)
    semestre = tabela.loc[tabela["ano"] == 2026].iloc[0]
    assert semestre["cobertura"] == "1º semestre"
    assert semestre["data_base"] == pd.Timestamp("2026-06-30")
    # Julho e agosto, 1% cada.
    assert semestre["fator"] == pytest.approx(1.01**2)
    assert semestre["itau_unibanco"] == SEMESTRE_2026["itau_unibanco"]
    assert semestre["itau"] is None or pd.isna(semestre["itau"])


def test_fusao_nao_duplica_2008_e_soma_exclui_2026():
    ipca = _ipca_sintetico()
    tabela = montar_tabela(ipca)
    ano_2007 = tabela.loc[tabela["ano"] == 2007].iloc[0]
    ano_2008 = tabela.loc[tabela["ano"] == 2008].iloc[0]
    assert ano_2007["grupo_itau"] == pytest.approx(ITAU[2007] + UNIBANCO[2007])
    assert ano_2008["itau"] is None or pd.isna(ano_2008["itau"])
    assert ano_2008["unibanco"] is None or pd.isna(ano_2008["unibanco"])
    assert ano_2008["grupo_itau"] == pytest.approx(ITAU_UNIBANCO[2008])
    cheios = tabela[tabela["cobertura"] == "Ano completo"]
    assert set(cheios["ano"]) == set(range(2002, 2026))
    assert 2026 not in set(cheios["ano"])
    # O lucro atualizado é nominal vezes o fator, sem outra transformação.
    assert ano_2007["bradesco_ipca"] == pytest.approx(BRADESCO[2007] * ano_2007["fator"])


def test_planilha_grava_os_nominais_oficiais(tmp_path: Path):
    ipca = _ipca_sintetico()
    tabela = montar_tabela(ipca)
    destino = gravar_planilha(tabela, ipca, tmp_path / "lucros.xlsx")
    wb = load_workbook(destino)
    ws = wb["Lucros_IPCA"]
    # Linha 5 = 2002. Colunas: D Itaú, E Unibanco, G Bradesco.
    assert ws["A5"].value == 2002
    assert ws["D5"].value == pytest.approx(2376.723)
    assert ws["E5"].value == pytest.approx(1010.0)
    assert ws["G5"].value == pytest.approx(2023.0)
    # 2008 é a linha 11. Itaú e Unibanco vazios; Itaú Unibanco = 7.803.
    assert ws["A11"].value == 2008
    assert ws["D11"].value is None
    assert ws["E11"].value is None
    assert ws["F11"].value == pytest.approx(7803.0)
    # 2025 é a linha 28. Controladora CVM.
    assert ws["A28"].value == 2025
    assert ws["F28"].value == pytest.approx(44857.0)
    assert ws["G28"].value == pytest.approx(23672.706)
    # Soma 2002–2025 não inclui 2026 (linha 29).
    assert ws["A29"].value == 2026
    assert ws["C29"].value == "1º semestre"
    assert ws["K30"].value == "=SUM(K5:K28)"
    assert "IPCA" in wb.sheetnames
    assert "Fontes_e_criterio" in wb.sheetnames
