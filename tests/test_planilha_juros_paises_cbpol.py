"""Taxas básicas diárias BIS: dias de semana no intervalo, uma aba por país."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from scripts.planilha_juros_paises_cbpol import (
    DATA_FIM,
    DATA_INICIO,
    classificar_colunas,
    extrair_series,
    gravar_planilha,
    nome_aba,
    parse_data_coluna,
    valor_taxa,
)


def test_parse_e_valor():
    assert parse_data_coluna("1995-01-02") == date(1995, 1, 2)
    assert parse_data_coluna("1995-01") is None
    assert parse_data_coluna("1995-13-40") is None
    assert valor_taxa(" 14.5 ") == 14.5
    assert valor_taxa("NaN") is None
    assert valor_taxa("") is None
    assert valor_taxa("0") == 0.0


def test_classificar_colunas_exclui_fim_de_semana_e_fora_do_periodo():
    header = [
        "FREQ",
        "1994-12-30",  # sexta, antes do período
        "1995-01-01",  # domingo
        "1995-01-02",  # segunda
        "1995-01-07",  # sábado
        "1995-01",
        "2026-08-31",  # segunda
        "2026-09-01",  # depois do período
    ]
    uteis, fins = classificar_colunas(header, DATA_INICIO, DATA_FIM)
    assert uteis == [(3, date(1995, 1, 2)), (6, date(2026, 8, 31))]
    assert fins == [2, 4]


def test_nome_aba_unico_e_curto():
    usados: set[str] = set()
    assert nome_aba("United States", usados) == "United States"
    assert nome_aba("Brasil/Sul", usados) == "BrasilSul"
    longo = nome_aba("X" * 40, usados)
    assert len(longo) == 31
    repetido = nome_aba("United States", usados)
    assert repetido != "United States"
    assert len(repetido) <= 31


def _csv_mini(path: Path) -> None:
    # 1995-01-01 domingo, 01-02 segunda, 01-07 sábado, 01-08 domingo, 2026-09-01 fora.
    linhas = [
        "FREQ,Frequency,REF_AREA,Reference area,TIME_FORMAT,Time Format,COMPILATION,"
        "1994-12-30,1995-01,1995-01-01,1995-01-02,1995-01-07,1995-01-08,2026-08-31,2026-09-01",
        "D,Daily,BR,Brazil,P1D,Daily,SELIC alvo,99,1,NaN,47.37,1.5,NaN,14,15",
        "D,Daily,US,United States,P1D,Daily,Fed funds,5.5,,5.5,5.5,5.5,,3.625,9",
        "M,Monthly,BR,Brazil,P1M,Monthly,mensal,1,1,1,1,1,1,1",
    ]
    path.write_text("\n".join(linhas) + "\n", encoding="utf-8")


def test_extrair_so_dias_de_semana_com_taxa(tmp_path: Path):
    origem = tmp_path / "mini.csv"
    _csv_mini(origem)
    series = extrair_series(origem)
    assert [s.codigo for s in series] == ["BR", "US"]

    brasil = series[0]
    assert brasil.observacoes == (
        (date(1995, 1, 2), 47.37),
        (date(2026, 8, 31), 14.0),
    )
    # sábado 1995-01-07 tem 1.5; domingo 01-08 é NaN e não entra na conta
    assert brasil.fins_de_semana_omitidos == 1
    assert brasil.definicao == "SELIC alvo"

    eua = series[1]
    assert eua.observacoes == (
        (date(1995, 1, 2), 5.5),
        (date(2026, 8, 31), 3.625),
    )
    # domingo 1995-01-01 = 5.5 e sábado 1995-01-07 = 5.5
    assert eua.fins_de_semana_omitidos == 2


def test_planilha_sem_sabado_domingo(tmp_path: Path):
    origem = tmp_path / "mini.csv"
    _csv_mini(origem)
    series = extrair_series(origem)
    destino = tmp_path / "juros.xlsx"
    abas = gravar_planilha(series, destino)

    wb = load_workbook(destino, read_only=True, data_only=True)
    assert wb.sheetnames == ["Indice", "Brazil", "United States"]
    assert abas == {"BR": "Brazil", "US": "United States"}

    indice = list(wb["Indice"].iter_rows(values_only=True))
    assert indice[1][0] == "Código"
    assert indice[2][0] == "BR"
    assert indice[2][3] == 2
    assert indice[2][8] == 1

    for nome in ("Brazil", "United States"):
        linhas = list(wb[nome].iter_rows(values_only=True))
        assert linhas[0] == ("Data", "Taxa básica (% a.a.)")
        datas = [linha[0].date() if hasattr(linha[0], "date") else linha[0] for linha in linhas[1:]]
        assert datas
        assert all(dt.weekday() < 5 for dt in datas)
        assert date(1995, 1, 1) not in datas
        assert date(1995, 1, 7) not in datas
        assert all(DATA_INICIO <= dt <= DATA_FIM for dt in datas)

    brasil = list(wb["Brazil"].iter_rows(values_only=True))
    assert brasil[1][1] == 47.37
    assert brasil[2][1] == 14
    wb.close()
